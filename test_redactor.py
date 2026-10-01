#!/usr/bin/env python3
"""
test_redactor.py - JARVIS escribe de verdad (redactor.py)
=========================================================
Ejecuta:  python test_redactor.py        (también con pytest)

Fija las cuatro cosas que hacían que pedirle un texto no sirviera de nada:

  * «escribe un ensayo…» tecleaba la frase literal en la ventana de JARVIS,
  * «abre el bloc de notas y escribe…» abría el Bloc de notas vacío,
  * el generador rellenaba con plantillas cuando el cerebro fallaba,
  * y una ruta relativa dejaba el documento dentro del proyecto.

Sin red y sin abrir nada: el cerebro es un doble que contesta según lo que se
le pide, y JARVIS_REDACTOR_ABRIR=0 impide lanzar programas.
"""
import json
import os
import sys
import tempfile

RAIZ = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, RAIZ)
_TMP = tempfile.mkdtemp(prefix="jarvis_redactor_")
os.environ["JARVIS_DOCUMENTOS"] = os.path.join(_TMP, "docs")
os.environ["JARVIS_REDACTOR_ESTADO"] = os.path.join(_TMP, "ultimo.json")
os.environ["JARVIS_REDACTOR_ABRIR"] = "0"
os.environ["JARVIS_FORMATO_TEXTO"] = "docx"

import redactor  # noqa: E402

# El diario de «deshaz eso» de las pruebas, aparte: si no, tras pasarlas un
# «deshaz eso» del señor intentaría revertir un documento de prueba.
try:
    import storage
    storage._almacen = storage.Storage(ruta=os.path.join(_TMP, "auditoria.db"),
                                       log=lambda *a: None)
except Exception:
    pass

fallos = []


# ── el doble del cerebro ─────────────────────────────────────────────────────
POEMA = ("# Oda a la Luna\n\nLuna de plata, farol del camino,\nque alumbras la senda "
         "del viejo molino.\n\nAñoro tu brillo cuando no estás,\ny cuento los días "
         "hasta verte jamás.")
CODIGO = ("Claro, aquí lo tiene:\n```python\nimport os\n\n\ndef main():\n    "
          "print('ordenando')\n\n\nif __name__ == '__main__':\n    main()\n```")
TABLA = {"titulo": "Gastos del mes", "hojas": [{
    "nombre": "Gastos", "columnas": ["Concepto", "Cantidad", "Precio", "Total"],
    "filas": [["Renta", 1, 8000, "=B2*C2"], ["Comida", 4, "1200.5", "=B3*C3"],
              ["Peligro", 1, 1, "=cmd|' /C calc'!A0"]], "totales": True}]}
DIAPOS = {"titulo": "Cambio climático", "subtitulo": "Causas y soluciones",
          "diapositivas": [{"titulo": "Qué es", "puntos": ["Más calor", "Más CO₂"],
                            "notas": "Explicar el concepto."},
                           {"titulo": "Soluciones", "puntos": ["Renovables"], "notas": ""}]}


class _Msg:
    def __init__(self, c):
        self.content = c


class _Eleccion:
    def __init__(self, c, fin="stop"):
        self.message = _Msg(c)
        self.finish_reason = fin


class _Resp:
    def __init__(self, c, fin="stop"):
        self.choices = [_Eleccion(c, fin)]
        self.usage = None


class _Completions:
    def __init__(self, cerebro):
        self.cerebro = cerebro

    def create(self, **kw):
        self.cerebro.llamadas.append(kw)
        if self.cerebro.falla:
            raise RuntimeError(self.cerebro.falla)
        sistema = kw["messages"][0]["content"]
        if self.cerebro.cortes:
            # Se queda sin tokens a mitad: el redactor debe pedirle que siga.
            self.cerebro.cortes -= 1
            return _Resp("# Ensayo largo\n\nPrimera mitad del texto", fin="length")
        if "corrigiendo un documento" in sistema:
            return _Resp("# Oda a la Luna\n\nLuna de plata.\n\nUna estrofa nueva al final.")
        if "programador" in sistema:
            return _Resp(CODIGO)
        if "hoja de cálculo" in sistema:
            return _Resp(json.dumps(TABLA))
        if "presentación" in sistema:
            return _Resp(json.dumps(DIAPOS))
        if kw["messages"][-1]["content"].startswith("Te has quedado cortado"):
            return _Resp(" y la segunda mitad, ya terminada.")
        return _Resp("¡Claro, señor! Aquí tiene:\n\n" + POEMA + "\n\n¿Desea que lo amplíe?")


class _Cliente:
    def __init__(self, cerebro):
        self.chat = type("Chat", (), {})()
        self.chat.completions = _Completions(cerebro)


class Cerebro:
    def __init__(self, falla="", cortes=0):
        self.falla, self.cortes, self.llamadas = falla, cortes, []


class Nucleo:
    """Lo mínimo del núcleo que usa el redactor."""
    nombre_agente = "JARVIS"
    _serie = [0]

    def __init__(self, *cerebros):
        Nucleo._serie[0] += 1
        self.serie = Nucleo._serie[0]
        self.cerebros = list(cerebros) or [Cerebro()]
        self.history = [{"role": "system", "content": "sistema"}]
        self.medios = []

    def _proveedores(self):
        # Nombres únicos por núcleo: cerebro_salud recuerda en memoria quién
        # falló, y la cuarentena de una prueba no debe vaciar la siguiente.
        # (Con id() no basta: Python reutiliza la dirección de un objeto libre.)
        return [(f"prueba-{self.serie}-{i}", f"http://127.0.0.1:{9000 + i}/v1", "modelo", "x")
                for i in range(len(self.cerebros))]

    def _cliente_llm(self, url, clave):
        return _Cliente(self.cerebros[int(url.split(":")[2][:4]) - 9000])

    def save_media_history(self, tipo, encargo, ruta):
        self.medios.append((tipo, ruta))

    def decir(self, frase):
        """Como el núcleo: atiende y apunta la respuesta en el historial."""
        r = redactor.atender_detalle(self, frase, log=lambda *a: None)
        if r:
            self.history += [{"role": "user", "content": frase},
                             {"role": "assistant", "content": r["mensaje"]}]
        return r


def _ruta(mensaje: str) -> str:
    for linea in mensaje.splitlines():
        if linea.startswith("Está en "):
            return linea[len("Está en "):].strip()
    return ""


# ── 1. qué entiende ──────────────────────────────────────────────────────────
ENCARGOS = [
    # (frase, modo, formato, app)
    ("abre el bloc de notas y escribe un poema sobre la luna", "generar", "txt", "notepad"),
    ("escríbeme en el blog de notas una lista de 10 películas de terror", "generar", "txt", "notepad"),
    ("escribe en el bloc de notas hola mundo", "literal", "txt", "notepad"),
    ("escribe en el bloc de notas que mañana compro pan", "literal", "txt", "notepad"),
    ("pon en el bloc de notas algo bonito para mi mamá", "generar", "txt", "notepad"),
    ("hazme un documento de Word sobre la Revolución Mexicana", "generar", "docx", "word"),
    ("abre Word y escribe una carta a mi jefe pidiendo vacaciones", "generar", "docx", "word"),
    ("crea un excel con mis gastos del mes", "generar", "xlsx", "excel"),
    ("hazme una presentación de PowerPoint sobre el cambio climático", "generar", "pptx", "powerpoint"),
    ("genera un PDF con un informe de ventas", "generar", "pdf", ""),
    ("redacta una carta de renuncia", "generar", "docx", ""),
    ("hazme un script en Python que ordene archivos por fecha", "generar", "py", ""),
    ("hazme una página web para mi taller mecánico", "generar", "html", ""),
    ("crea un archivo de texto con un resumen de la segunda guerra mundial", "generar", "txt", ""),
    ("genera un archivo csv con los países de América y su capital", "generar", "csv", ""),
    ("escribe aquí un poema", "generar", "txt", ""),
    ("pásalo al bloc de notas", "anterior", "txt", "notepad"),
    ("guarda eso en un word", "anterior", "docx", "word"),
    ("¿puedes escribirme un poema?", "generar", "txt", ""),
    ("necesito que me redactes un oficio para la escuela", "generar", "docx", ""),
    ("crea un archivo python que ordene mis descargas por tipo", "generar", "py", ""),
    ("usa word y escríbeme una carta para el banco", "generar", "docx", "word"),
    ("hazme un plan de entrenamiento en excel", "generar", "xlsx", "excel"),
    ("dame un resumen en pdf de la segunda guerra mundial", "generar", "pdf", ""),
    ("traduce al inglés lo que me dijiste y pásalo a word", "generar", "docx", "word"),
]

# Lo que NO es suyo: otros módulos lo atienden o es conversación.
AJENAS = [
    "escribe hola mundo",                      # tecleo de pc_control
    "escríbele a Juan que llego tarde",        # mensajería (WhatsApp)
    "escribe un correo a Luis que vemos mañana",
    "manda un whatsapp a mamá diciendo que llego tarde",
    "anota que mañana tengo dentista",         # notas de JARVIS
    "crea una nota con la lista del super",
    "¿cómo se escribe excepción?",
    "crea un archivo llamado test.txt en el escritorio",   # jarvis_skills._archivos
    "hazme un resumen de la fotosíntesis",     # charla
    "cuéntame un chiste", "escribe un chiste",
    "hazme un video para mi presentación", "hazme una imagen de un gato",
    "hazme un informe de la práctica 3",       # informe.py (plugin estudiar)
    "hazme una tabla periódica",               # ciencias
    "abre el bloc de notas", "abre el bloc de notas y escribe",
    "¿qué puedo hacer con excel?", "describe esta imagen", "transcribe este audio",
    "dame la receta del pozole", "apaga el pc y abre el bloc de notas",
    # investigar y guardar: jarvis_skills._investigar, que cita sus fuentes
    "busca información sobre la luna y guárdala en un documento",
    "investiga sobre los volcanes y guárdalo en un archivo html",
    "crea una carpeta llamada proyectos en el escritorio", "crea un recordatorio para mañana",
    "pon la canción de despacito", "pásame el archivo al móvil", "hazme una lista de reproducción",
    "escribe en Word lo que te voy a dictar",
]


def test_entiende_los_encargos():
    redactor.olvidar_ultimo()
    for frase, modo, formato, app in ENCARGOS:
        e = redactor.entender(frase)
        assert e, f"«{frase}» no se reconoció como encargo"
        assert (e["modo"], e["formato"], e["app"]) == (modo, formato, app), \
            f"«{frase}» -> {(e['modo'], e['formato'], e['app'])}, esperaba {(modo, formato, app)}"


def test_no_secuestra_lo_ajeno():
    redactor.olvidar_ultimo()
    for frase in AJENAS:
        assert redactor.entender(frase) is None, f"«{frase}» no es un encargo de escribir"


def test_destino_nombre_y_extension():
    e = redactor.entender("escribe un ensayo de 1000 palabras sobre la inteligencia "
                          "artificial y guárdalo en el escritorio")
    assert e["palabras"] == 1000
    assert e["ubicacion"] == redactor.carpeta("escritorio")
    assert e["tema"] == "la inteligencia artificial"
    e = redactor.entender("crea un documento llamado Ideas de negocio con 10 ideas de negocio")
    assert e["nombre"] == "Ideas de negocio"
    e = redactor.entender("escribe un poema y guárdalo en la carpeta Poemas")
    assert e["ubicacion"] == os.path.join(redactor.carpeta_base(), "Poemas")
    # La literal conserva tildes y mayúsculas aunque las regex vean texto plano.
    e = redactor.entender("en el bloc de notas escribe: Mañana compro PAN y leche")
    assert e["contenido"] == "Mañana compro PAN y leche"


# ── 2. escribe de verdad, en cada formato ────────────────────────────────────
def test_escribe_cada_formato():
    redactor.olvidar_ultimo()
    nucleo = Nucleo()

    txt = _ruta(nucleo.decir("abre el bloc de notas y escribe un poema sobre la luna")["mensaje"])
    assert txt.endswith(".txt") and os.path.isfile(txt)
    with open(txt, "rb") as f:
        crudo = f.read()
    assert crudo.startswith(b"\xef\xbb\xbf"), "el .txt lleva BOM para el Bloc de notas antiguo"
    texto = crudo.decode("utf-8-sig")
    assert "Añoro tu brillo" in texto and "farol del camino,\nque alumbras" in texto
    assert "Aquí tiene" not in texto and "Desea" not in texto, "se coló la cháchara del modelo"

    docx = _ruta(nucleo.decir("hazme un documento de Word sobre la luna")["mensaje"])
    from docx import Document
    parrafos = [(p.style.name, p.text) for p in Document(docx).paragraphs]
    assert ("Title", "Oda a la Luna") in parrafos
    assert any("Añoro tu brillo" in t for _s, t in parrafos)

    pdf = _ruta(nucleo.decir("genera un PDF con un poema sobre la luna")["mensaje"])
    with open(pdf, "rb") as f:
        assert f.read(5) == b"%PDF-"

    xlsx = _ruta(nucleo.decir("crea un excel con mis gastos del mes")["mensaje"])
    from openpyxl import load_workbook
    hoja = load_workbook(xlsx)["Gastos"]
    assert hoja["A1"].value == "Concepto" and hoja.freeze_panes == "A2"
    assert hoja["D2"].value == "=B2*C2", "las fórmulas se quedan como fórmulas"
    assert hoja["C3"].value == 1200.5, "los números llegan como números"
    assert str(hoja["D4"].value).startswith("'"), "una fórmula DDE se neutraliza"
    assert hoja["A5"].value == "Total" and hoja["C5"].value == "=SUM(C2:C4)"

    pptx = _ruta(nucleo.decir("hazme una presentación sobre el cambio climático")["mensaje"])
    from pptx import Presentation
    diapos = Presentation(pptx).slides
    assert len(diapos) == 3
    assert diapos[1].notes_slide.notes_text_frame.text == "Explicar el concepto."

    py = _ruta(nucleo.decir("hazme un script en Python que ordene archivos por fecha")["mensaje"])
    assert os.path.basename(py) == "ordene_archivos_por_fecha.py"
    with open(py, encoding="utf-8") as f:
        codigo = f.read()
    assert codigo.startswith("import os") and "```" not in codigo
    compile(codigo, py, "exec")

    assert len(nucleo.medios) == 6, "cada documento queda en el historial de medios"


def test_lo_dictado_y_lo_de_antes():
    redactor.olvidar_ultimo()
    nucleo = Nucleo()
    ruta = _ruta(nucleo.decir("escribe en el bloc de notas que mañana compro pan")["mensaje"])
    assert os.path.basename(ruta) == "Mañana compro pan.txt"
    with open(ruta, encoding="utf-8-sig") as f:
        assert f.read().strip() == "Mañana compro pan"
    assert not nucleo.cerebros[0].llamadas, "lo dictado no gasta cerebro"

    nucleo.history.append({"role": "assistant", "content": "La fotosíntesis convierte la luz "
                                                           "en energía química."})
    ruta = _ruta(nucleo.decir("guarda eso en un word")["mensaje"])
    from docx import Document
    assert "fotosíntesis" in Document(ruta).paragraphs[0].text

    # «Tradúcelo y pásalo a Word»: lo de antes es el MATERIAL, no el resultado.
    nucleo.history.append({"role": "assistant", "content": "El sol es una estrella."})
    nucleo.decir("traduce al inglés lo que me dijiste y pásalo a word")
    peticion = nucleo.cerebros[0].llamadas[-1]["messages"][1]["content"]
    assert "El sol es una estrella." in peticion and "traduce al inglés" in peticion


# ── 3. el cerebro: relevo, continuación y honestidad ─────────────────────────
def test_relevo_entre_cerebros():
    redactor.olvidar_ultimo()
    sin_saldo, bueno = Cerebro(falla="401 credit balance too low"), Cerebro()
    nucleo = Nucleo(sin_saldo, bueno)
    r = nucleo.decir("escríbeme un poema de amor")
    assert os.path.isfile(_ruta(r["mensaje"])), r["mensaje"]
    assert sin_saldo.llamadas and bueno.llamadas, "si uno no tiene saldo, escribe el siguiente"


def test_sigue_si_se_corta():
    nucleo = Nucleo(Cerebro(cortes=1))
    texto = redactor.preguntar(nucleo, "sistema", "usuario", log=lambda *a: None)
    assert texto.endswith("la segunda mitad, ya terminada."), texto


def test_sin_cerebro_no_hay_plantillas():
    redactor.olvidar_ultimo()
    carpeta = redactor.carpeta_base()
    antes = set(os.listdir(carpeta)) if os.path.isdir(carpeta) else set()
    r = Nucleo(Cerebro(falla="Connection refused")).decir("escribe un ensayo sobre la luna")
    assert r and "No he podido escribirlo" in r["mensaje"]
    despues = set(os.listdir(carpeta)) if os.path.isdir(carpeta) else set()
    assert despues == antes, "sin cerebro no se crea ningún documento de relleno"


# ── 4. el último documento ───────────────────────────────────────────────────
def test_seguimientos():
    redactor.olvidar_ultimo()
    nucleo = Nucleo()
    primera = _ruta(nucleo.decir("abre el bloc de notas y escribe un poema sobre la luna")["mensaje"])

    r = nucleo.decir("hazlo más largo")
    assert r and "He cambiado" in r["mensaje"]
    with open(primera, encoding="utf-8-sig") as f:
        assert "Una estrofa nueva" in f.read()

    r = nucleo.decir("pásalo a pdf")
    pdf = _ruta(r["mensaje"])
    assert pdf.endswith(".pdf") and os.path.isfile(pdf)

    assert "está en" in nucleo.decir("¿dónde está el documento?")["mensaje"]
    # Lo que no habla del documento no lo toca.
    for frase in ("lee el correo", "pon la lista de reproducción", "abre el archivo de configuración",
                  "hazme otro documento sobre los gatos"):
        e = redactor.entender(frase, core=nucleo)
        assert not e or e.get("modo") != "seguimiento", f"«{frase}» no va con el último documento"


def test_nunca_sobrescribe():
    carpeta = os.path.join(_TMP, "colision")
    os.makedirs(carpeta, exist_ok=True)
    with open(os.path.join(carpeta, "Carta.txt"), "w") as f:
        f.write("original")
    assert redactor.ruta_libre(carpeta, "Carta.txt").endswith("Carta (2).txt")
    assert redactor.nombre_archivo('Informe: "ventas" 2026?', "docx") == "Informe ventas 2026.docx"
    assert redactor.nombre_archivo("con", "txt") == "con_.txt", "CON es un nombre reservado"


def test_no_ejecuta_lo_que_escribe():
    # Fuera de Windows no hay editor que lanzar: se dice, y NUNCA se pasa un
    # programa a xdg-open (que lo ejecutaría con su intérprete).
    os.environ["JARVIS_REDACTOR_ABRIR"] = "1"
    try:
        ruta = os.path.join(_TMP, "peligro.py")
        with open(ruta, "w") as f:
            f.write("raise SystemExit('no me ejecutes')\n")
        if os.name != "nt":
            abierto, motivo = redactor.abrir(ruta, "")
            assert not abierto and "no sé abrirlo" in motivo
    finally:
        os.environ["JARVIS_REDACTOR_ABRIR"] = "0"


# ── 5. los otros caminos ─────────────────────────────────────────────────────
def test_tecleo_sin_tildes_perdidas():
    assert redactor.texto_a_teclear("teclea Mañana voy a la piñata") == "Mañana voy a la piñata"
    assert redactor.texto_a_teclear("escribe aquí: «Buenos días»") == "Buenos días"
    assert redactor.texto_a_teclear("escribe a Juan que llego tarde") == "", "eso es mensajería"
    assert redactor.es_ventana_de_jarvis("J.A.R.V.I.S. · Origen - Microsoft Edge")
    assert redactor.es_ventana_de_jarvis("Administrador: C:\\WINDOWS\\system32\\cmd.exe - python jarvis.py")
    assert not redactor.es_ventana_de_jarvis("Documento1 - Word")


def test_pc_control_no_teclea_describe():
    with open(os.path.join(RAIZ, "pc_control.py"), encoding="utf-8") as f:
        fuente = f.read()
    assert "pyautogui.typewrite(mm.group(1)" not in fuente, "el tecleo literal sin tildes sigue ahí"
    assert r'\b(?:escribe|teclea)\s+\S' in fuente, "sin \\b, «describe…» se teclea"
    assert "redactor.teclear(" in fuente


def test_red_del_chat():
    """El modelo contesta con [OPEN:bloc de notas] y el texto: no se abre vacío."""
    redactor.olvidar_ultimo()
    abiertos = []
    nucleo = Nucleo()
    nucleo._open_app = abiertos.append
    nota = redactor.guardar_respuesta(
        nucleo, "[OPEN:bloc de notas] Claro, señor. Aquí tiene su lista: leche, huevos, pan, "
                "fruta, verdura, arroz, frijoles y café.", "bloc de notas", log=lambda *a: None)
    assert nota and not abiertos
    with open(redactor.ultimo()["ruta"], encoding="utf-8-sig") as f:
        assert f.read().startswith("leche, huevos")
    # Sin texto que guardar, se abre como siempre.
    assert redactor.guardar_respuesta(nucleo, "[OPEN:bloc de notas] Abriendo el Bloc de notas, señor.",
                                      "bloc de notas", log=lambda *a: None) == ""
    assert abiertos == ["bloc de notas"]
    assert redactor.app_de_escritura("Bloc de notas") == "notepad"
    assert redactor.app_de_escritura("spotify") == ""


def test_herramienta_del_modelo():
    import permisos
    from herramientas_llm import Herramientas
    caja = Herramientas(None, log=lambda *a: None)
    nombres = [d["function"]["name"] for d in caja.definiciones()]
    assert "crear_documento" in nombres and hasattr(caja, "_t_crear_documento")
    assert permisos.evaluar("crear_documento") == "directo", "crear un documento nuevo no pide «confirma»"
    redactor.olvidar_ultimo()
    caja.core = Nucleo()
    msg = caja._t_crear_documento({"encargo": "un poema sobre la luna", "abrir_con": "bloc de notas"})
    assert os.path.isfile(_ruta(msg)), msg
    msg = caja._t_crear_documento({"encargo": "nota", "formato": "txt", "nombre": "compras",
                                   "contenido": "pan y leche"})
    assert os.path.basename(_ruta(msg)) == "compras.txt"
    # Un nombre suelto de documento ya no acaba dentro del proyecto.
    destino = redactor.ruta_documento("poema.txt")
    assert os.path.dirname(destino) == redactor.carpeta_base()


def test_generador_de_la_web():
    import jarvis_generator
    gen = jarvis_generator.JarvisGenerator(log=lambda *a: None, core=Nucleo())
    r = gen.gen_word("generar word: carta formal")
    assert r["type"] == "word" and os.path.isfile(r["path"]) and not r.get("plantilla")
    from docx import Document
    assert any("Añoro" in p.text for p in Document(r["path"]).paragraphs), "contenido de verdad"
    gen = jarvis_generator.JarvisGenerator(log=lambda *a: None,
                                          core=Nucleo(Cerebro(falla="sin red")))
    r = gen.gen_word("generar word: carta formal")
    assert r.get("plantilla") is True, "si cae a la plantilla, se dice"


# ── ejecución sin pytest ─────────────────────────────────────────────────────
def main():
    pruebas = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    for prueba in pruebas:
        try:
            prueba()
            print(f"  [OK]  {prueba.__name__}")
        except AssertionError as e:
            fallos.append(f"{prueba.__name__}: {e}")
            print(f"  [MAL] {prueba.__name__}: {e}")
        except Exception as e:
            fallos.append(f"{prueba.__name__}: {type(e).__name__}: {e}")
            print(f"  [MAL] {prueba.__name__}: {type(e).__name__}: {e}")
    print()
    if fallos:
        print(f"FALLARON {len(fallos)} de {len(pruebas)}")
        return 1
    print(f"TODO EN VERDE ({len(pruebas)} pruebas)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
