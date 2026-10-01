#!/usr/bin/env python3
"""
redactor.py - Que JARVIS escriba de verdad
==========================================
Pedirle un texto a JARVIS no servía de nada, y no por un fallo sino por cuatro
que se tapaban unos a otros:

* «escribe un ensayo sobre X» acababa en el tecleo de pc_control, que mandaba
  LITERALMENTE «un ensayo sobre x» —sin tildes ni eñes— a la ventana que
  tuviera el foco. Casi siempre la del propio JARVIS: escribía en sí mismo.
* «abre el bloc de notas y escribe…» abría el Bloc de notas VACÍO. Ni las
  habilidades ni las herramientas del modelo sabían escribir dentro, y el
  texto, si llegaba, se quedaba en el chat recortado a tres frases.
* El generador le pedía el contenido a un proveedor fijo que se quedó sin
  saldo y, al fallar, rellenaba el documento con plantillas («Este documento
  aborda el tema de…»). Parecía que escribía; no escribía nada.
* La herramienta de escribir archivos pedía «confirma» siempre, y una ruta
  relativa acababa dentro de la carpeta del proyecto.

Aquí está el camino entero, de la frase al documento abierto con el texto:

    entender(frase)      qué texto, en qué formato, dónde y con qué programa
    preguntar(core, …)   el contenido, con la cadena de cerebros del núcleo:
                         si uno falla o se queda sin saldo, contesta el siguiente
    escribir(…)          .txt .md .docx .rtf .pdf .xlsx .csv .pptx, páginas
                         web, código y JSON, comprobado en disco
    abrir(ruta, app)     Bloc de notas, Word, Excel, PowerPoint, VS Code o el
                         navegador, ya con el texto dentro

Tres promesas:

* **Nunca una plantilla vacía.** Si ningún cerebro contesta, se dice y no se
  crea nada: un documento de relleno es peor que ninguno.
* **Nunca sobrescribe.** Si el nombre existe, se numera. Todo lo que crea queda
  en el diario de «deshaz eso».
* **Nunca ejecuta lo que escribe.** Un .py, un .bat o un .js se abren en un
  editor, no con su programa por defecto, que los ejecutaría.

Y recuerda el último documento, para «añádele un párrafo sobre…», «hazlo más
largo», «pásalo a PDF», «ábrelo» o «mándamelo al móvil».

Variables de entorno:

    JARVIS_DOCUMENTOS          carpeta de salida (por defecto Documentos\\JARVIS)
    JARVIS_REDACTOR_TIMEOUT    segundos por llamada al cerebro (300)
    JARVIS_REDACTOR_ABRIR=0    no abrir nada al terminar (pruebas, servidores)
    JARVIS_FORMATO_TEXTO       formato de los textos formales sin formato
                               pedido: docx (si hay Word), txt o md
"""
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import unicodedata

_HOME = os.path.expanduser("~")
_PREFS = os.path.join(_HOME, "Descargas", "JARVIS", "Prefs")
_SIN_VENTANA = 0x08000000 if os.name == "nt" else 0


def _entero(nombre: str, por_defecto: int) -> int:
    """Entero del entorno; vacío o mal escrito cuenta como no puesto."""
    try:
        return int((os.getenv(nombre) or "").strip() or por_defecto)
    except ValueError:
        return por_defecto


# Un ensayo largo con el cerebro de casa tarda minutos, no segundos: con el
# timeout de 60 s del chat la respuesta se perdía entera a medio escribir.
_TIMEOUT = _entero("JARVIS_REDACTOR_TIMEOUT", 300)
# «hazlo más largo» o «ábrelo» se refieren al documento si es así de reciente.
_RECIENTE_S = _entero("JARVIS_REDACTOR_RECIENTE", 20 * 60)


# ══════════════════════════════════════════════════════════════════════════
# TEXTO PLANO PARA LAS REGEX
# ══════════════════════════════════════════════════════════════════════════
def _plano(texto: str) -> str:
    """Minúsculas y sin tildes, con la MISMA longitud que el original.

    Las regex se pasan sobre esta versión, pero el contenido se recorta del
    original con los mismos índices: así «escribe en el bloc de notas:
    Mañana compro pan» deja «Mañana compro pan» con su eñe y su mayúscula.
    """
    salida = []
    for c in texto or "":
        base = unicodedata.normalize("NFD", c)[:1] or c
        bajo = base.lower()
        salida.append(bajo[:1] if bajo else base)
    return "".join(salida)


def _espacios(texto: str) -> str:
    return re.sub(r"[ \t]{2,}", " ", texto or "").strip()


def _tapar(plano: str, patrones) -> str:
    """Cambia por espacios lo que case: el texto conserva su longitud.

    Sirve para buscar el género sin mirar el destino: en «escribe en el bloc
    de NOTAS…» la palabra «notas» es el programa, no lo que se pide escribir.
    """
    trozos = list(plano)
    for rx in patrones:
        for m in rx.finditer(plano):
            for i in range(m.start(), m.end()):
                trozos[i] = " "
    return "".join(trozos)


def _quitar_spans(orig: str, plano: str, patrones) -> tuple:
    """Quita de los dos textos (alineados) lo que case con cualquier patrón."""
    tramos = []
    for rx in patrones:
        tramos += [m.span() for m in rx.finditer(plano) if m.end() > m.start()]
    if not tramos:
        return orig, plano
    tramos.sort()
    unidos = [list(tramos[0])]
    for ini, fin in tramos[1:]:
        if ini <= unidos[-1][1]:
            unidos[-1][1] = max(fin, unidos[-1][1])
        else:
            unidos.append([ini, fin])
    o, p, previo = [], [], 0
    for ini, fin in unidos:
        o.append(orig[previo:ini])
        p.append(plano[previo:ini])
        previo = fin
    o.append(orig[previo:])
    p.append(plano[previo:])
    # El mismo colapso en los dos: los espacios están en los mismos sitios,
    # así que siguen alineados.
    return _espacios("".join(o)), _espacios("".join(p))


# ══════════════════════════════════════════════════════════════════════════
# CARPETAS DE VERDAD (Escritorio y Documentos pueden vivir en OneDrive)
# ══════════════════════════════════════════════════════════════════════════
# Con OneDrive, el Escritorio que ve el señor es C:\Users\X\OneDrive\Escritorio
# y no C:\Users\X\Desktop. Escribir en el segundo es dejar el archivo donde
# nadie lo va a ver nunca. Windows sabe cuál es el bueno: se le pregunta.
_CARPETAS = {
    "escritorio": ("B4BFCC3A-DB2C-424C-B029-7FE99A87C641", ("Desktop", "Escritorio")),
    "documentos": ("FDD39AD0-238F-46AF-ADB4-6C85480369C7", ("Documents", "Documentos")),
    # «Descargas» a secas NO: ~/Descargas es la carpeta de datos de JARVIS.
    "descargas": ("374DE290-123F-4565-9164-39C4925E467B", ("Downloads",)),
    "imagenes": ("33E28130-4E1E-4676-835A-98395C3BC3BB", ("Pictures", "Imágenes")),
    "musica": ("4BD8D571-6D19-48D3-BE97-422220080E43", ("Music", "Música")),
    "videos": ("18989B1D-99B5-455B-841C-AB7C74E4DDFC", ("Videos", "Vídeos")),
}
_ALIAS_CARPETA = {
    "desktop": "escritorio", "documents": "documentos", "mis documentos": "documentos",
    "downloads": "descargas", "fotos": "imagenes", "pictures": "imagenes",
    "music": "musica",
}


def _carpeta_windows(guid: str) -> str:
    import ctypes
    import uuid
    ident = (ctypes.c_ubyte * 16).from_buffer_copy(uuid.UUID(guid).bytes_le)
    ptr = ctypes.c_void_p()
    hr = ctypes.windll.shell32.SHGetKnownFolderPath(ctypes.byref(ident), 0, None,
                                                    ctypes.byref(ptr))
    if hr != 0 or not ptr.value:
        return ""
    try:
        return ctypes.wstring_at(ptr.value)
    finally:
        ctypes.windll.ole32.CoTaskMemFree(ptr)


def carpeta(nombre: str) -> str:
    """Ruta real de Escritorio, Documentos, Descargas... ('' si no se conoce)."""
    clave = _plano(nombre or "").strip()
    clave = _ALIAS_CARPETA.get(clave, clave)
    if clave not in _CARPETAS:
        return ""
    guid, nombres = _CARPETAS[clave]
    if os.name == "nt":
        try:
            ruta = _carpeta_windows(guid)
            if ruta and os.path.isdir(ruta):
                return ruta
        except Exception:
            pass
    for n in nombres:
        ruta = os.path.join(_HOME, n)
        if os.path.isdir(ruta):
            return ruta
    return os.path.join(_HOME, nombres[0])


def carpeta_base() -> str:
    """Donde van los documentos si el señor no dice otra cosa."""
    propia = (os.getenv("JARVIS_DOCUMENTOS") or "").strip()
    if propia:
        return os.path.abspath(os.path.expanduser(propia))
    return os.path.join(carpeta("documentos"), "JARVIS")


# ══════════════════════════════════════════════════════════════════════════
# FORMATOS Y PROGRAMAS
# ══════════════════════════════════════════════════════════════════════════
_FAMILIA = {"txt": "texto", "md": "texto", "rtf": "texto", "docx": "texto",
            "pdf": "texto", "xlsx": "tabla", "csv": "tabla", "pptx": "diapos"}

# Extensión -> cómo se le nombra el lenguaje al modelo.
_LENGUAJES = {
    "py": "Python", "js": "JavaScript", "ts": "TypeScript", "java": "Java",
    "cpp": "C++", "c": "C", "cs": "C#", "php": "PHP", "rb": "Ruby", "go": "Go",
    "rs": "Rust", "kt": "Kotlin", "swift": "Swift", "sql": "SQL", "sh": "Bash",
    "ps1": "PowerShell", "bat": "un archivo por lotes de Windows (.bat)",
    "ino": "Arduino (C++)", "html": "HTML con el CSS y el JavaScript dentro",
    "css": "CSS", "xml": "XML", "yml": "YAML", "json": "JSON", "tex": "LaTeX",
}

_NOMBRE_APP = {
    "notepad": "el Bloc de notas", "word": "Word", "excel": "Excel",
    "powerpoint": "PowerPoint", "wordpad": "WordPad",
    "vscode": "Visual Studio Code", "navegador": "el navegador",
}
_EXES = {"word": "winword.exe", "excel": "excel.exe", "powerpoint": "powerpnt.exe",
         "wordpad": "wordpad.exe"}
_APP_DE_FORMATO = {"docx": "word", "rtf": "word", "xlsx": "excel", "csv": "excel",
                   "pptx": "powerpoint", "txt": "notepad", "md": "notepad",
                   "html": "navegador"}
_FORMATO_DE_APP = {"notepad": "txt", "word": "docx", "wordpad": "rtf",
                   "excel": "xlsx", "powerpoint": "pptx"}


def familia(formato: str) -> str:
    formato = (formato or "").lower().lstrip(".")
    if formato in _FAMILIA:
        return _FAMILIA[formato]
    return "codigo" if formato in _LENGUAJES else "texto"


_APPS_CACHE = {}


def _ruta_app(exe: str) -> str:
    """Ruta de Word, Excel o PowerPoint por «App Paths»: no suelen estar en el
    PATH, y un «start winword» sin Word instalado saca una ventana de error."""
    if exe in _APPS_CACHE:
        return _APPS_CACHE[exe]
    ruta = ""
    if os.name == "nt":
        try:
            import winreg
            clave = r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths" + "\\" + exe
            for raiz in (winreg.HKEY_CURRENT_USER, winreg.HKEY_LOCAL_MACHINE):
                for vista in (0, winreg.KEY_WOW64_64KEY, winreg.KEY_WOW64_32KEY):
                    try:
                        with winreg.OpenKey(raiz, clave, 0, winreg.KEY_READ | vista) as k:
                            valor = os.path.expandvars(str(winreg.QueryValueEx(k, "")[0]).strip('"'))
                        if valor and os.path.isfile(valor):
                            ruta = valor
                            break
                    except OSError:
                        continue
                if ruta:
                    break
        except Exception:
            pass
    ruta = ruta or shutil.which(exe) or ""
    _APPS_CACHE[exe] = ruta
    return ruta


def _ruta_vscode() -> str:
    if "vscode" in _APPS_CACHE:
        return _APPS_CACHE["vscode"]
    ruta = ""
    code = shutil.which("code")
    if code:
        # bin\code.cmd -> Code.exe: abrir el .exe no levanta una consola.
        exe = os.path.normpath(os.path.join(os.path.dirname(code), "..", "Code.exe"))
        ruta = exe if os.path.isfile(exe) else code
    _APPS_CACHE["vscode"] = ruta
    return ruta


def hay_word() -> bool:
    return bool(_ruta_app("winword.exe"))


# ══════════════════════════════════════════════════════════════════════════
# ENTENDER LA ORDEN
# ══════════════════════════════════════════════════════════════════════════
# Todo sobre texto plano (sin tildes, en minúsculas).
_RX_V_ESCRIBIR = re.compile(
    r"\b(?:escrib(?:e|i|ir|a|as|an|amos|iria|irias)"
    r"(?:me|te|le|les|nos|lo|la|los|las|melo|mela|selo|sela)?"
    r"|redact(?:a|ar|e|es|en)(?:me|te|le|les|nos|lo|la|los|las|melo|mela|selo|sela)?"
    r"|compon(?:er)?(?:me|le|lo|la|melo|mela)?|componga|compongas)\b")
_RX_V_CREAR = re.compile(
    r"\b(?:haz|hazme|hazlo|hazla|hazmelo|hazmela|hacer|hacerme|hacerlo|hacermelo"
    r"|haga|hagame|hagalo|hagamelo|hagas|haces|harias|haria"
    r"|cre(?:a|ar|e|es)(?:me|le|lo|la|melo|mela|selo)?"
    r"|gener(?:a|ar|e|es)(?:me|le|lo|la|melo|mela)?"
    r"|prepar(?:a|ar|e|es)(?:me|le|lo|la|melo|mela)?"
    r"|elabor(?:a|ar|e|es)(?:me|le|lo|la|melo|mela)?"
    r"|arm(?:ame|amelo|amela|ar|arme)"
    r"|dame|deme|quiero|quisiera|necesito|ocupo|me gustaria)\b")
_RX_V_PONER = re.compile(
    r"\b(?:pon|ponme|ponlo|ponla|ponmelo|ponmela|ponle|poner|ponerlo|ponerla"
    r"|ponga|pongame|pongalo|pongala"
    r"|pasa|pasame|pasalo|pasala|pasamelo|pasamela|pasar|pasarlo|pasarla"
    r"|guarda|guardame|guardalo|guardala|guardamelo|guardamela|guardar|guardarlo"
    r"|guardarla|guarde|guardelo"
    r"|copia|copiame|copialo|copiala|copiamelo|copiamela"
    r"|exporta|exportalo|exportala|exportamelo|exportamela"
    r"|mete|metelo|metela|metemelo|vuelca|vuelcalo|vuelcala"
    r"|apuntalo|apuntala|apuntamelo|anotalo|anotala|anotamelo)\b")
# Verbo con «lo/la» pegado y nada más: «pásalo al Bloc de notas» habla de lo
# que JARVIS acaba de decir.
_RX_CLITICO = re.compile(
    r"\b(?:escrib\w*?(?:lo|la)|redact\w*?(?:lo|la)|pon(?:me)?(?:lo|la)"
    r"|pas(?:a|ame|ar)(?:lo|la)|guard(?:a|ame|ar|e)(?:lo|la)|copi(?:a|ame)(?:lo|la)"
    r"|export(?:a|ame)(?:lo|la)|met(?:e|eme)(?:lo|la)|vuelca(?:lo|la)"
    r"|apunta(?:me)?(?:lo|la)|anota(?:me)?(?:lo|la))\b")

_APPS = (r"(?:(?:bloc|block|blog) de notas?|notepad|wordpad|word|excel|power ?point"
         r"|visual studio code|vs ?code)")
_FMTS = (r"(?:pdf|docx|txt|xlsx|csv|pptx|markdown|json|html|rtf|archivo de texto"
         r"|texto plano|hoja de calculo|hoja de excel)")
# Dónde lo quiere: el programa o el formato, en una construcción que lo pida
# («en el bloc de notas», «un excel», «pásalo a pdf», «documento de word»).
# «un tutorial de excel» NO: ahí Excel es el tema, no el destino.
_RX_OBJETIVO = re.compile(
    r"\b(?:abre|abreme|abrir|abrelo|abrela|inicia|usa|usando|utiliza|en|con|al|a"
    r"|desde|dentro de)\s+(?:el\s+|la\s+|un\s+|una\s+|mi\s+)?" + _APPS + r"\b"
    r"|\b(?:un|una|el|la|al|a|como|formato|tipo|en)\s+" + _FMTS + r"\b"
    r"|\b(?:documento|archivo|fichero|hoja|libro|presentacion|plantilla)\s+"
    r"(?:de\s+|en\s+)?(?:" + _APPS + "|" + _FMTS + r")\b"
    r"|\b(?:un|una)\s+(?:word|excel|power ?point)\b")
_RX_CONTENEDOR = re.compile(
    r"\b(?:en|a|como|dentro de)\s+(?:un|una|el|la|mi)\s+(?:documento|archivo|fichero)\b"
    r"|\b(?:un|una)\s+(?:documento|archivo|fichero)\s+(?:de|con|sobre|que|para|acerca"
    r"|donde|en|llamado|llamada|titulado|titulada)\b")

_PISTAS_APP = (
    (re.compile(r"\b(?:bloc|block|blog) de notas?\b|\bnotepad\b"), "notepad"),
    (re.compile(r"\bwordpad\b"), "wordpad"),
    (re.compile(r"\bword\b"), "word"),
    (re.compile(r"\bexcel\b|hoja de excel"), "excel"),
    (re.compile(r"\bpower ?point\b"), "powerpoint"),
    (re.compile(r"\bvisual studio code\b|\bvs ?code\b"), "vscode"),
)
_PISTAS_FORMATO = (
    (re.compile(r"\bpdf\b"), "pdf"),
    (re.compile(r"\bdocx\b|\bword\b"), "docx"),
    (re.compile(r"\brtf\b|\bwordpad\b"), "rtf"),
    (re.compile(r"\bxlsx\b|\bexcel\b|hoja de calculo"), "xlsx"),
    (re.compile(r"\bcsv\b"), "csv"),
    (re.compile(r"\bpptx\b|\bpower ?point\b"), "pptx"),
    (re.compile(r"\bmarkdown\b"), "md"),
    (re.compile(r"\bjson\b"), "json"),
    (re.compile(r"\bhtml\b"), "html"),
    (re.compile(r"\btxt\b|archivo de texto|texto plano|(?:bloc|block|blog) de notas?|\bnotepad\b"), "txt"),
)
_PISTAS_LENGUAJE = (
    (re.compile(r"\bpython\b|\.py\b"), "py"),
    (re.compile(r"\btypescript\b|\.ts\b"), "ts"),
    (re.compile(r"\bjavascript\b|\bnode(?:\.?js)?\b|\.js\b|\bjs\b"), "js"),
    (re.compile(r"\bjava\b"), "java"),
    (re.compile(r"\bc\+\+|\bcpp\b"), "cpp"),
    (re.compile(r"\bc#|\bcsharp\b|\bc sharp\b"), "cs"),
    (re.compile(r"\bphp\b"), "php"),
    (re.compile(r"\bruby\b"), "rb"),
    (re.compile(r"\bgolang\b|\ben go\b"), "go"),
    (re.compile(r"\brust\b"), "rs"),
    (re.compile(r"\bkotlin\b"), "kt"),
    (re.compile(r"\bswift\b"), "swift"),
    (re.compile(r"\bsql\b"), "sql"),
    (re.compile(r"\bbash\b|\bshell script\b|\.sh\b"), "sh"),
    (re.compile(r"\bpowershell\b|\.ps1\b"), "ps1"),
    (re.compile(r"\bbatch\b|\.bat\b|\barchivo bat\b|\.cmd\b"), "bat"),
    (re.compile(r"\barduino\b|\.ino\b"), "ino"),
    (re.compile(r"\bhtml\b"), "html"),
    (re.compile(r"\bcss\b"), "css"),
    (re.compile(r"\bxml\b"), "xml"),
    (re.compile(r"\byaml\b|\byml\b"), "yml"),
    (re.compile(r"\blatex\b|\.tex\b"), "tex"),
    (re.compile(r"\ben c\b|\blenguaje c\b"), "c"),
)

# Qué se le pide. Fuerza: «fuerte» basta con cualquier verbo de crear
# («hazme un ensayo»); «débil» necesita «escribe/redacta» o un destino
# explícito («hazme un resumen» se contesta charlando; «escríbeme un resumen»
# o «un resumen en Word» no); «corto» (un chiste, una frase) solo con destino.
# El último valor es el formato si el señor no dice ninguno.
_GENEROS = [
    # código
    (r"\bpaginas? web\b|\bsitios? web\b|\bpagina de internet\b|\blanding(?: page)?\b",
     "página web", "fuerte", "html"),
    (r"\bscripts?\b", "script", "fuerte", "py"),
    (r"\b(?:archivo|fichero)s?\s+(?:de\s+|en\s+)?(?:python|javascript|typescript|java|c\+\+|c#"
     r"|php|ruby|golang|rust|kotlin|swift|sql|bash|powershell|batch|bat|arduino|css)\b",
     "código", "fuerte", "py"),
    (r"\bcodigos?\b(?! postal| de barras| qr| civil| penal)", "código", "fuerte", "py"),
    (r"\bprogramas?\b", "programa", "debil", "py"),
    (r"\bfunciones\b|\bfuncion\b|\bclases? en\b|\balgoritmos?\b", "código", "debil", "py"),
    # hojas de cálculo
    (r"\bhojas? de calculo\b", "hoja de cálculo", "fuerte", "xlsx"),
    (r"\binventarios?\b", "inventario", "fuerte", "xlsx"),
    (r"\bpresupuestos?\b", "presupuesto", "fuerte", "xlsx"),
    (r"\bfacturas?\b", "factura", "fuerte", "xlsx"),
    (r"\bcotizacion(?:es)?\b", "cotización", "fuerte", "xlsx"),
    (r"\bnominas?\b", "nómina", "fuerte", "xlsx"),
    (r"\b(?:control|registro|lista) de gastos\b", "control de gastos", "fuerte", "xlsx"),
    (r"\btablas?\b(?! de multiplicar| periodica| de verdad)", "tabla", "fuerte", "xlsx"),
    (r"\bhorarios?\b|\bcronogramas?\b", "horario", "debil", "xlsx"),
    # presentaciones (una «carta de presentación» es una carta)
    (r"(?<!carta de )(?<!video de )(?<!mensaje de )(?<!texto de )(?<!pagina de )"
     r"\bpresentacion(?:es)?\b|\bdiapositivas?\b|\bslides?\b", "presentación", "fuerte", "pptx"),
    # prosa formal
    (r"\bensayos?\b", "ensayo", "fuerte", "docx"),
    (r"\binformes?\b", "informe", "fuerte", "docx"),
    (r"\breportes?\b", "reporte", "fuerte", "docx"),
    (r"\bmonografias?\b", "monografía", "fuerte", "docx"),
    (r"\btesis\b|\btesinas?\b", "tesis", "fuerte", "docx"),
    (r"\barticulos?\b", "artículo", "fuerte", "docx"),
    (r"\btrabajos? (?:escrito|de investigacion|final|practico|academico)s?\b"
     r"|\btrabajo sobre\b", "trabajo", "fuerte", "docx"),
    (r"\bcontratos?\b", "contrato", "fuerte", "docx"),
    (r"\bcurricul(?:um|o)s?\b|\bcv\b|\bhoja de vida\b", "currículum", "fuerte", "docx"),
    (r"\bcartas? (?:formal|de (?:renuncia|presentacion|recomendacion|motivos|motivacion"
     r"|reclamacion|reclamo|queja|disculpa|agradecimiento|despido|aceptacion|poder"
     r"|referencia))", "carta formal", "fuerte", "docx"),
    (r"\bsolicitud(?:es)?\b", "solicitud", "fuerte", "docx"),
    (r"\boficios?\b", "oficio", "fuerte", "docx"),
    (r"\bmemorandums?\b|\bmemos?\b", "memorándum", "fuerte", "docx"),
    (r"\bactas?\b", "acta", "fuerte", "docx"),
    (r"\bpropuestas?\b", "propuesta", "fuerte", "docx"),
    (r"\bmanual(?:es)?\b", "manual", "fuerte", "docx"),
    (r"\bguias?\b", "guía", "fuerte", "docx"),
    (r"\btutorial(?:es)?\b", "tutorial", "fuerte", "docx"),
    (r"\bexamen(?:es)?\b|\bcuestionarios?\b|\bencuestas?\b", "cuestionario", "fuerte", "docx"),
    (r"\btemarios?\b", "temario", "fuerte", "docx"),
    (r"\bplan de (?:negocios?|empresa|marketing)\b", "plan de negocio", "fuerte", "docx"),
    (r"\bbiografias?\b", "biografía", "fuerte", "docx"),
    (r"\bresenas?\b|\bcriticas?\b", "reseña", "fuerte", "docx"),
    (r"\bcronicas?\b", "crónica", "fuerte", "docx"),
    (r"\bdocumentos?\b", "documento", "fuerte", "docx"),
    # prosa ligera: al Bloc de notas
    (r"\bcartas?\b", "carta", "fuerte", "txt"),
    (r"\bpoemas?\b|\bpoesias?\b|\bversos?\b|\bhaikus?\b|\bsonetos?\b", "poema", "fuerte", "txt"),
    (r"\bcuentos?\b|\brelatos?\b|\bfabulas?\b|\bleyendas?\b|\bnovelas?\b|\bcapitulos?\b"
     r"|\bmicrorrelatos?\b", "cuento", "fuerte", "txt"),
    (r"\bcancion(?:es)?\b|\bletras? de (?:una |la )?cancion\b|\bcorridos?\b", "canción", "fuerte", "txt"),
    (r"\bguion(?:es)?\b|\bobras? de teatro\b|\bdialogos?\b|\bmonologos?\b", "guion", "fuerte", "txt"),
    (r"\bdiscursos?\b|\bbrindis\b|\bpalabras de (?:bienvenida|despedida|agradecimiento)\b",
     "discurso", "fuerte", "txt"),
    (r"\binvitacion(?:es)?\b|\bfelicitacion(?:es)?\b|\bdedicatorias?\b", "invitación", "fuerte", "txt"),
    # débiles
    (r"\bresumen(?:es)?\b|\bsintesis\b", "resumen", "debil", "txt"),
    (r"\blistas?\b|\bchecklist\b", "lista", "debil", "txt"),
    (r"\bplan(?:es)?\b|\bplanificacion\b|\bplaneacion\b", "plan", "debil", "txt"),
    (r"\bideas\b|\bconsejos\b|\bpasos\b|\binstrucciones\b|\brazones\b|\bpreguntas\b",
     "lista", "debil", "txt"),
    (r"\bmensajes?\b|\bcorreos?\b|\be-?mails?\b|\bmails?\b|\bposts?\b"
     r"|\bpublicacion(?:es)?\b|\btuits?\b", "mensaje", "debil", "txt"),
    (r"\bnotas?\b", "nota", "debil", "txt"),
    (r"\btextos?\b|\bparrafos?\b|\bredaccion(?:es)?\b|\bcomposicion(?:es)?\b|\bescritos?\b",
     "texto", "debil", "txt"),
    (r"\bexplicacion(?:es)?\b|\bdescripcion(?:es)?\b|\bdefinicion(?:es)?\b",
     "explicación", "debil", "txt"),
    (r"\brecetas?\b|\bdietas?\b|\bmenus?\b|\brutinas?\b", "receta", "debil", "txt"),
    (r"\bhistorias?\b|\bdiarios?\b", "historia", "debil", "txt"),
    (r"\bapuntes\b|\besquemas?\b|\bglosarios?\b|\bfichas?\b", "apuntes", "debil", "txt"),
    # cortos: al chat salvo destino explícito
    (r"\bchistes?\b|\badivinanzas?\b|\btrabalenguas\b|\bfrases?\b|\brefran(?:es)?\b"
     r"|\bpiropos?\b|\bsaludos?\b", "frases", "corto", "txt"),
]
_GENEROS = [(re.compile(rx), tipo, fuerza, defecto) for rx, tipo, fuerza, defecto in _GENEROS]

# Lo que se pide crear y NO es texto: lo atienden otros módulos.
_MEDIOS = {"imagen", "imagenes", "foto", "fotos", "dibujo", "dibujos", "ilustracion",
           "logo", "logos", "video", "videos", "gif", "modelo", "holograma", "melodia",
           "musica", "audio", "sonido", "grafica", "grafico", "diagrama", "plano",
           "mapa", "alarma", "recordatorio", "temporizador", "carpeta", "habilidad",
           "copia", "captura", "backup", "respaldo", "cita", "evento", "reunion"}

_RX_EXCLUIR = re.compile(
    # preguntas sobre cómo escribir, no encargos
    r"\bcomo se (?:escribe|redacta|hace|dice|pronuncia)\b"
    r"|^\W*(?:que|como|cuando|donde|por que|cual|cuales|quien|puedo|se puede)\b"
    r"(?!.*\b(?:puedes|podrias|me (?:haces|harias|escribes|redactas|creas|preparas)|"
    r"quiero|necesito)\b)"
    # lo atienden módulos propios
    r"|\binforme (?:de|sobre|para) (?:la |las |mis |una |el )?(?:practica|laboratorio)"
    r"|\bmemoria de practicas\b|\bborrador (?:de|para) (?:la |una )?entrega\b"
    r"|\bdict(?:ado|ar|e|ame)\b|\bhabilidad(?:es)?\b"
    # investigar y guardar: jarvis_skills._investigar, con Wikipedia de fuente
    r"|\binvestiga(?:me|r|lo|la)?\b|\bhaz(?:me)? una investigacion\b"
    r"|\bbusca(?:me)?\s+(?:informacion|info|datos)\b|\baverigua\w*"
    r"|\bescribete\b|\bmodo invitado\b")

_RX_MENSAJERIA = re.compile(
    r"\b(?:whatsapp|wsp|wasap|telegram|sms|instagram|facebook|messenger|twitter|gmail|outlook)\b"
    r"|\b(?:manda|mandale|mandales|envia|enviale|enviales|responde|respondele|contesta"
    r"|contestale)\b"
    r"|\bescrib(?:e|ir)?les?\s+(?:a|al)\s+\w"
    r"|\b(?:correo|email|mail|mensaje)\s+(?:a|al|para)\s+[\w ]{1,40}?\s+(?:que|diciendo)\b"
    r"|\b(?:correo|email|mail|mensaje)\s+(?:a|al|para)\s+[\w ]{1,40}?:")

_RX_AQUI = re.compile(
    r"\b(?:aqui|aca|en la ventana(?: activa| abierta)?|donde (?:esta|tengo|tenga) el cursor"
    r"|en lo que tengo abierto|en el documento abierto|en este documento|en esta ventana)\b")
_RX_LITERAL = re.compile(
    r"\b(?:que diga|que dice|que ponga|que pone|que contenga|diciendo"
    r"|con el (?:siguiente )?(?:texto|contenido|mensaje)|con (?:este|el) texto"
    r"|lo siguiente|textualmente|literalmente|tal cual)\b\s*[:,]?\s*")
_RX_ANTERIOR = re.compile(
    r"\b(?:eso|esto|lo anterior|lo de antes|todo eso|todo esto|ese texto|este texto"
    r"|el texto anterior|tu (?:ultima )?respuesta|la respuesta|lo que (?:me )?(?:dijiste"
    r"|has dicho|acabas de decir|escribiste|explicaste|contaste|respondiste"
    r"|contestaste))\b")

_RX_UBICACION = re.compile(
    r"\s*\b(?:y\s+)?(?:(?:guarda|guardalo|guardala|guardamelo|ponlo|ponla|dejalo"
    r"|dejala)\s+)?(?:en|a|al|dentro de|hacia)\s+(?:el\s+|la\s+|los\s+|las\s+|mi\s+|mis\s+)?"
    r"(?:carpeta\s+(?:de\s+)?)?(?P<u>escritorio|desktop|mis documentos|documentos"
    r"|descargas|downloads|imagenes|fotos|musica|videos)\b")
_RX_RUTA = re.compile(r"\s*\b(?:en|a)\s+(?P<r>[a-z]:[\\/][^\s\"'«»]*)")
_RX_CARPETA = re.compile(
    r"\s*\b(?:y\s+)?(?:guardalo\s+|guardala\s+|ponlo\s+|ponla\s+)?(?:en|a|dentro de)\s+"
    r"(?:la|una|mi)\s+carpeta\s+(?:llamada\s+|de nombre\s+)?[\"'«]?(?P<c>[\w][\w .-]{0,40}?)"
    r"[\"'»]?(?=\s+(?:del|de|en|y|con|sobre|que)\b|[,.;]|\s*$)")
_RX_NOMBRE = re.compile(
    r"\s*\b(?:llamad[oa]|que se llame|que se llama|con (?:el )?nombre(?: de)?|de nombre"
    r"|titulad[oa]|con (?:el )?titulo(?: de)?)\s+[\"'«]?(?P<n>.+?)[\"'»]?"
    r"(?=\s+(?:en|con|y|sobre|para|que|acerca|de \d)\b|[,;]|\s*$)")
_RX_ABRIR_APP = re.compile(
    r"\b(?:abre|abreme|abrir|inicia|lanza|ejecuta)\s+(?:el\s+|la\s+|un\s+|una\s+|mi\s+)?"
    + _APPS + r"\s*(?:,\s*)?(?:y\s+)?")
_RX_RELLENO = re.compile(r"\bpor favor\b|\bjarvis\b|\bultron\b|^\s*oye\b|^\s*ey\b")
_RX_DESTINO = re.compile(
    r"\s*\b(?:y\s+)?(?:abre(?:lo|la|melo|mela)?|abrir|en|con|al|a|desde|dentro de|usando"
    r"|como)\s+(?:el\s+|la\s+|un\s+|una\s+|mi\s+)?(?:" + _APPS + "|" + _FMTS + r")\b"
    r"(?:\s+(?:de\s+)?(?:" + _APPS + "|" + _FMTS + r"))?"
    r"|\s*\b(?:en|a|como|dentro de)\s+(?:un|una|el|la|mi)\s+(?:documento|archivo|fichero)"
    r"(?:\s+(?:de\s+|en\s+)?(?:" + _APPS + "|" + _FMTS + r"))?\b"
    r"|\s*\by\s+(?:abre|abrelo|abrela|abremelo|abremela|muestramelo|muestramela)\b")


def _primero(rx, texto: str):
    return rx.search(texto)


def _detectar_destino(p: str) -> tuple:
    """(app, formato) pedidos explícitamente, o ('', '')."""
    app = formato = ""
    for m in _RX_OBJETIVO.finditer(p):
        trozo = m.group(0)
        if not app:
            for rx, nombre in _PISTAS_APP:
                if rx.search(trozo):
                    app = nombre
                    break
        if not formato:
            for rx, nombre in _PISTAS_FORMATO:
                if rx.search(trozo):
                    formato = nombre
                    break
    return app, formato


def _detectar_lenguaje(p: str) -> str:
    for rx, ext in _PISTAS_LENGUAJE:
        if rx.search(p):
            return ext
    return ""


def _detectar_genero(p: str, desde: int = 0):
    """El género que aparece ANTES tras el verbo: en «un ensayo sobre el código
    de Hammurabi» manda «ensayo»; «código» es el tema."""
    mejor = None
    for rx, tipo, fuerza, defecto in _GENEROS:
        m = rx.search(p, desde)
        if m and (mejor is None or m.start() < mejor[0]):
            mejor = (m.start(), tipo, fuerza, defecto, m.end())
    if mejor is None and desde:
        return _detectar_genero(p, 0)
    if mejor is None:
        return None
    return {"tipo": mejor[1], "fuerza": mejor[2], "defecto": mejor[3],
            "inicio": mejor[0], "fin": mejor[4]}


def _objeto_creado(p: str, desde: int) -> str:
    """La primera palabra con contenido tras el verbo: «hazme UN VIDEO para mi
    presentación» pide un vídeo, no una presentación."""
    resto = p[desde:]
    for palabra in re.findall(r"[a-z0-9]+", resto)[:5]:
        if palabra in ("me", "te", "le", "un", "una", "unos", "unas", "el", "la", "los",
                       "las", "mi", "mis", "otro", "otra", "nuevo", "nueva", "buen",
                       "buena", "pequeno", "pequena", "corto", "corta", "breve", "lo",
                       "por", "favor", "rapido", "rapidamente", "aqui", "porfa"):
            continue
        return palabra
    return ""


def _palabras(p: str, tipo: str, familia_fmt: str) -> int:
    """Cuánto escribir: lo que diga el señor o lo razonable para el género."""
    m = re.search(r"\b(\d{2,5})\s*palabras\b", p)
    if m:
        return max(30, min(int(m.group(1)), 12000))
    numeros = {"una": 1, "dos": 2, "tres": 3, "cuatro": 4, "cinco": 5, "seis": 6,
               "siete": 7, "ocho": 8, "nueve": 9, "diez": 10}
    m = re.search(r"\b(\d{1,2}|una|dos|tres|cuatro|cinco|seis|siete|ocho|nueve|diez)\s+"
                  r"(?:paginas|pagina|cuartillas|cuartilla|hojas|hoja|carillas)\b", p)
    if m:
        n = numeros.get(m.group(1)) or int(m.group(1))
        return max(200, min(n * 400, 12000))
    base = {
        "poema": 180, "canción": 260, "carta": 350, "carta formal": 400, "mensaje": 160,
        "nota": 120, "lista": 300, "frases": 150, "invitación": 150, "discurso": 700,
        "cuento": 900, "guion": 900, "resumen": 450, "explicación": 600, "receta": 450,
        "plan": 700, "texto": 500, "historia": 900, "apuntes": 900, "ensayo": 1200,
        "informe": 1300, "reporte": 1100, "monografía": 2500, "tesis": 3000,
        "artículo": 1100, "trabajo": 1800, "contrato": 1200, "currículum": 500,
        "solicitud": 350, "oficio": 350, "memorándum": 350, "acta": 500,
        "propuesta": 1000, "manual": 1800, "guía": 1400, "tutorial": 1400,
        "cuestionario": 800, "temario": 600, "plan de negocio": 2000, "biografía": 1100,
        "reseña": 700, "crónica": 900, "documento": 1000,
    }.get(tipo, 600 if familia_fmt == "texto" else 0)
    if re.search(r"\b(?:corto|corta|breve|pequeno|pequena|cortito|cortita|rapido|sencillo"
                 r"|sencilla)\b", p):
        base = int(base * 0.5)
    elif re.search(r"\b(?:largo|larga|extenso|extensa|detallado|detallada|completo"
                   r"|completa|a fondo|exhaustivo|exhaustiva|profundo|profunda)\b", p):
        base = int(base * 2)
    return base


def _ubicacion(orig: str, p: str) -> str:
    """Carpeta pedida, o '' para la de siempre."""
    m = _RX_RUTA.search(p)
    if m:
        return orig[m.start("r"):m.end("r")].rstrip(".,;")
    m = _RX_UBICACION.search(p)
    base = carpeta(m.group("u")) if m else ""
    m_c = _RX_CARPETA.search(p)
    if m_c:
        nombre = orig[m_c.start("c"):m_c.end("c")].strip()
        if _plano(nombre) not in _CARPETAS and _plano(nombre) not in _ALIAS_CARPETA:
            if base:
                return os.path.join(base, nombre)
            # Una carpeta que ya existe en un sitio conocido gana a crear otra.
            for sitio in ("escritorio", "documentos", "descargas"):
                for real in _listar(carpeta(sitio)):
                    if _plano(real) == _plano(nombre) and os.path.isdir(os.path.join(carpeta(sitio), real)):
                        return os.path.join(carpeta(sitio), real)
            return os.path.join(carpeta_base(), nombre)
    return base


def _listar(ruta: str) -> list:
    try:
        return os.listdir(ruta)
    except Exception:
        return []


def _nombre_pedido(orig: str, p: str) -> str:
    m = _RX_NOMBRE.search(p)
    if not m:
        return ""
    return orig[m.start("n"):m.end("n")].strip().strip("\"'«».,;")


_CANTIDADES = {"un", "una", "unos", "unas", "varios", "varias", "algunos", "algunas", "dos",
               "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve", "diez", "veinte",
               "cien", "alguna", "muchos", "muchas"}


def _descripcion(orig: str, p: str, genero) -> str:
    """Lo que se pide, tal como lo dijo: «poema sobre la luna», «carta de
    renuncia», «10 chistes», «script en Python que ordene archivos». Para el
    aviso."""
    if not genero:
        return ""
    inicio = genero["inicio"]
    previo = re.search(r"\b(\d+|" + "|".join(sorted(_CANTIDADES - {"un", "una"})) + r")\s+$",
                       p[:inicio])
    if previo:
        inicio = previo.start()
    texto = orig[inicio:]
    plano = p[inicio:]
    texto, plano = _quitar_spans(texto, plano, (_RX_DESTINO, _RX_UBICACION, _RX_CARPETA,
                                                _RX_NOMBRE, _RX_RUTA, _RX_RELLENO))
    corte = re.search(r"[.;:!?¿¡\n]|,\s*(?:y\s+)?(?:abre|guarda|pon|ponlo|ponla)", texto)
    if corte:
        texto = texto[:corte.start()]
    return " ".join(texto.split()[:14]).strip(" ,-")


def _tema(orig: str, p: str, genero) -> str:
    """El tema, para el nombre del archivo: lo que va tras «sobre…» o, si no
    lo hay, la descripción entera."""
    m = re.search(r"\b(?:sobre|acerca de|de como|respecto a|relacionad[oa] con|hablando de)\s+(.+)$", p)
    if m:
        texto, plano = orig[m.start(1):], p[m.start(1):]
        texto, plano = _quitar_spans(texto, plano, (_RX_DESTINO, _RX_UBICACION, _RX_CARPETA,
                                                    _RX_NOMBRE, _RX_RUTA, _RX_RELLENO))
        texto = re.split(r"[.;:!?¿¡\n]|,", texto)[0]
        return " ".join(texto.split()[:9]).strip(" -")
    return " ".join(_descripcion(orig, p, genero).split()[:9])


def _limpiar_encargo(orig: str) -> str:
    """La frase para el cerebro, sin «abre el bloc de notas y» ni la carpeta."""
    p = _plano(orig)
    o, p = _quitar_spans(orig, p, (_RX_ABRIR_APP, _RX_RELLENO, _RX_UBICACION, _RX_CARPETA,
                                   _RX_RUTA))
    o = re.sub(r"^[\s,.;:]+|[\s,;:]+$", "", o)
    return o or orig.strip()


def _ultima_respuesta(core) -> str:
    for m in reversed(list(getattr(core, "history", None) or [])):
        if isinstance(m, dict) and m.get("role") == "assistant":
            texto = str(m.get("content") or "").strip()
            if texto:
                return texto
    return ""


def _resto_tras_verbo(orig: str, p: str, verbo) -> tuple:
    """Lo que viene después del verbo, sin destino, carpeta ni nombre."""
    ini = verbo.end()
    o, pp = orig[ini:], p[ini:]
    o, pp = _quitar_spans(o, pp, (_RX_DESTINO, _RX_UBICACION, _RX_CARPETA, _RX_NOMBRE,
                                  _RX_RUTA, _RX_AQUI, _RX_RELLENO))
    m = re.match(r"^(?:me\s+|nos\s+)?(?:lo siguiente|esto|esta frase|este texto)?\s*[:,]?\s*", pp)
    if m and m.end():
        o, pp = o[m.end():], pp[m.end():]
    return o.strip(" ,;"), pp.strip(" ,;")


def _dictado(orig: str, p: str, verbo) -> str:
    """El texto exacto si el señor lo dicta con marca: entre comillas, tras
    «que diga» o «con el texto», o tras dos puntos. '' si no hay marca."""
    m = re.search(r"[«\"“]([^»\"”]{1,6000})[»\"”]", orig)
    if m:
        return m.group(1).strip()
    m = _RX_LITERAL.search(p, verbo.start())
    if m:
        o, pp = orig[m.end():], p[m.end():]
        o, pp = _quitar_spans(o, pp, (_RX_DESTINO, _RX_UBICACION, _RX_CARPETA, _RX_RUTA))
        return o.strip(" ,;")
    m = re.search(r":\s*(.+)$", orig[verbo.end():], re.S)
    if m:
        return m.group(1).strip()
    return ""


def _mayuscula(texto: str) -> str:
    return texto[:1].upper() + texto[1:] if texto else ""


def entender(texto: str, core=None):
    """¿Esto es un encargo de escribir? Devuelve el encargo (dict) o None.

    No llama a ningún cerebro: solo regex, para poder ir delante de todo sin
    costar nada a las frases que no son suyas.
    """
    if not texto or not str(texto).strip() or len(str(texto)) > 6000:
        return None
    orig = str(texto).strip()
    p = _plano(orig)

    seguimiento = _entender_seguimiento(orig, p, core)
    if seguimiento:
        return seguimiento

    if _RX_EXCLUIR.search(p):
        return None
    v_esc = _RX_V_ESCRIBIR.search(p)
    v_crear = _RX_V_CREAR.search(p)
    v_poner = _RX_V_PONER.search(p)
    verbos = [v for v in (v_esc, v_crear, v_poner) if v]
    if not verbos:
        return None
    verbo = min(verbos, key=lambda m: m.start())

    app, formato = _detectar_destino(p)
    contenedor = bool(_RX_CONTENEDOR.search(p))
    objetivo = bool(app or formato or contenedor)
    # El género se busca SIN el destino ni el nombre: en «escribe en el bloc
    # de notas…» o «un archivo de texto…», «notas» y «texto» no son el encargo.
    p_gen = _tapar(p, (_RX_OBJETIVO, _RX_DESTINO, _RX_CONTENEDOR, _RX_NOMBRE, _RX_CARPETA))
    genero = _detectar_genero(p_gen, verbo.end())
    lenguaje = _detectar_lenguaje(p)
    pegar = bool(v_esc and re.search(_RX_V_ESCRIBIR.pattern + r"\s+(?:me\s+)?"
                                     + _RX_AQUI.pattern, p))

    if _RX_MENSAJERIA.search(p) and not app:
        return None
    if _objeto_creado(p_gen, verbo.end()) in _MEDIOS:
        return None
    if genero and genero["tipo"] == "nota" and not (app or formato):
        # «crea una nota», «escribe una nota que diga…»: el archivo de notas de
        # JARVIS (jarvis_skills._notas), no un documento.
        return None

    # ¿Es un encargo?
    encargo = False
    if v_esc and (genero or objetivo or pegar):
        encargo = not (genero and genero["fuerza"] == "corto" and not objetivo)
    if not encargo and (v_crear or v_esc) and genero:
        if genero["fuerza"] == "fuerte" or objetivo:
            encargo = True
        elif genero["defecto"] in _LENGUAJES and lenguaje:
            encargo = True       # «hazme un programa en Python que…»
    if not encargo and v_crear and objetivo and re.search(
            r"\b(?:sobre|acerca|de|del|con|para|que|donde|explicando)\s+\w", p[verbo.end():]):
        encargo = True           # «hazme un Excel con mis gastos»
    if not encargo and v_poner and (objetivo or pegar):
        encargo = True           # «pásalo al Bloc de notas»
    if not encargo:
        return None

    # ¿Qué contenido? Dictado, lo de antes o redactado.
    v_contenido = v_esc or v_poner
    literal = _dictado(orig, p, v_contenido or verbo) if (v_contenido or not genero) else ""
    anterior = False
    sin_genero = False
    if not literal and v_contenido and not genero:
        resto_o, resto_p = _resto_tras_verbo(orig, p, v_contenido)
        if not resto_o:
            # «pásalo al Bloc de notas», «guárdamelo en un Word»
            anterior = bool(_RX_CLITICO.search(p) or v_poner)
        elif _RX_ANTERIOR.match(resto_p) and len(resto_p.split()) <= 7:
            anterior = True      # «guarda eso en un Word», «pon lo que me dijiste»
        elif re.match(r"(?:algo|alguna cosa|unas? |unos |vari[oa]s |\d+ |dos |tres |cinco |diez )",
                      resto_p):
            # «pon en el bloc de notas algo bonito para mi madre»: se pide que
            # lo escriba él, no que copie la frase.
            genero = {"tipo": "texto", "fuerza": "debil", "defecto": "txt", "inicio": 0, "fin": 0}
            sin_genero = True
        else:
            literal = _mayuscula(resto_o[4:].strip() if resto_p.startswith("que ") else resto_o)
    usar_anterior = False
    if anterior and re.search(r"\b(?:traduce|traducelo|traducela|resume|resumelo|resumela"
                              r"|corrige|corrigelo|mejora|mejoralo|reescribe|reescribelo|amplia"
                              r"|amplialo|explica|explicalo|ordena|ordenalo|organiza|organizalo"
                              r"|pasa a limpio|simplifica|simplificalo|adapta|adaptalo)\b",
                              p[:v_contenido.start()]):
        # Lo de antes es el material; el encargo, lo que se pide hacer con él.
        anterior, usar_anterior = False, True
    if v_crear and not v_contenido and literal:
        # «crea un archivo que diga hola» ya lo hace jarvis_skills._archivos,
        # con su nombre y su carpeta de siempre.
        return None
    if not literal and not anterior and not genero and not (v_crear and objetivo) \
            and not usar_anterior:
        return None

    # Formato y programa.
    fam_genero = familia(genero["defecto"]) if genero else ""
    if not formato and app:
        formato = _FORMATO_DE_APP.get(app, "")
    nombre = _nombre_pedido(orig, p)
    ext_nombre = os.path.splitext(nombre)[1].lower().lstrip(".")
    if ext_nombre and (ext_nombre in _FAMILIA or ext_nombre in _LENGUAJES):
        formato = ext_nombre
    if genero and fam_genero == "codigo":
        if formato in ("", "txt") or app in ("notepad", "vscode"):
            formato = lenguaje or genero["defecto"]
    elif not formato and lenguaje and app == "vscode":
        formato = lenguaje
    if not formato:
        if genero:
            formato = genero["defecto"]
            # «Un documento con diez ideas»: quien dice documento piensa en Word.
            if formato == "docx" or (formato == "txt" and re.search(r"\bdocumentos?\b", p)):
                formato = _formato_texto_formal()
        elif v_crear and contenedor:
            formato = _formato_texto_formal()
        else:
            formato = "txt"

    tipo = genero["tipo"] if genero else ("documento" if familia(formato) == "texto"
                                          else {"tabla": "hoja de cálculo", "diapos": "presentación",
                                                "codigo": "código"}[familia(formato)])
    descripcion = "" if sin_genero else _descripcion(orig, p, genero)
    if not genero or sin_genero:
        # Sin género, lo que sigue al destino describe el encargo: «un Excel
        # CON MIS GASTOS DEL MES».
        resto_o, resto_p = _resto_tras_verbo(orig, p, verbo)
        resto_o, resto_p = _quitar_spans(resto_o, resto_p, (_RX_CONTENEDOR, _RX_OBJETIVO))
        art = re.match(r"(?:un|una|el|la)\s+(?=(?:con|de|del|sobre|para|acerca|que)\b)", resto_p)
        if art:
            resto_o, resto_p = resto_o[art.end():], resto_p[art.end():]
        if re.match(r"(?:algo|alguna cosa|unas?|unos|varios|varias|\d+)\b", resto_p):
            descripcion = resto_o.strip()
        elif re.match(r"(?:con|de|del|sobre|para|acerca|que)\b", resto_p):
            descripcion = f"{tipo} {resto_o}".strip()
    return {
        "modo": "literal" if literal else ("anterior" if anterior else "generar"),
        "frase": orig,
        "encargo": _limpiar_encargo(orig),
        "contenido": literal,
        "tipo": tipo,
        "formato": formato,
        "app": app,
        "ubicacion": _ubicacion(orig, p),
        "nombre": nombre,
        "tema": _tema(orig, p, None if sin_genero else genero),
        "descripcion": descripcion,
        "palabras": _palabras(p, tipo, familia(formato)),
        "pegar": pegar,
        "usar_anterior": usar_anterior,
    }


def _formato_texto_formal() -> str:
    """Un ensayo o un informe sin formato pedido: Word si lo hay."""
    pedido = (os.getenv("JARVIS_FORMATO_TEXTO") or "").strip().lower().lstrip(".")
    if pedido in ("docx", "txt", "md", "rtf", "pdf"):
        return pedido
    return "docx" if (os.name != "nt" or hay_word()) else "txt"


def es_encargo(texto: str, core=None) -> bool:
    try:
        return entender(texto, core=core) is not None
    except Exception:
        return False


# ══════════════════════════════════════════════════════════════════════════
# EL ÚLTIMO DOCUMENTO
# ══════════════════════════════════════════════════════════════════════════
_ULTIMO = {}
_CANDADO = threading.Lock()


def _ruta_estado() -> str:
    return (os.getenv("JARVIS_REDACTOR_ESTADO") or "").strip() or \
        os.path.join(_PREFS, "ultimo_documento.json")


def _recordar(info: dict):
    with _CANDADO:
        _ULTIMO.clear()
        _ULTIMO.update(info)
        try:
            ruta = _ruta_estado()
            os.makedirs(os.path.dirname(ruta), exist_ok=True)
            copia = dict(info)
            copia["contenido"] = str(copia.get("contenido") or "")[:200000]
            with open(ruta, "w", encoding="utf-8") as f:
                json.dump(copia, f, ensure_ascii=False, indent=1)
        except Exception:
            pass


def ultimo() -> dict:
    """El último documento que escribió JARVIS ({} si ninguno)."""
    with _CANDADO:
        if _ULTIMO:
            return dict(_ULTIMO)
    try:
        with open(_ruta_estado(), encoding="utf-8") as f:
            datos = json.load(f) or {}
        if datos.get("ruta"):
            with _CANDADO:
                _ULTIMO.update(datos)
            return dict(datos)
    except Exception:
        pass
    return {}


def olvidar_ultimo():
    with _CANDADO:
        _ULTIMO.clear()
    try:
        os.remove(_ruta_estado())
    except OSError:
        pass


def _fue_lo_ultimo(core, ult: dict) -> bool:
    """¿Lo último que dijo JARVIS fue entregar ese documento?"""
    if core is None:
        return True
    previa = _ultima_respuesta(core).strip()
    if not previa:
        return False
    # La respuesta tal cual o, si el núcleo le añadió algo, la que nombra el archivo.
    return previa == str(ult.get("respuesta") or "").strip() or \
        bool(ult.get("ruta") and ult["ruta"] in previa)


_GENERICOS = r"(?:documento|archivo|fichero|texto|escrito)"
# Cómo se puede llamar el último documento según lo que era.
_SINONIMOS = {
    "carta formal": "carta", "poema": "poema|poesia|verso|versos", "canción": "cancion|letra",
    "cuento": "cuento|relato|historia", "presentación": "presentacion|diapositivas|power ?point",
    "hoja de cálculo": "hoja|tabla|excel", "tabla": "tabla|hoja|excel", "inventario": "inventario|tabla|hoja|excel",
    "presupuesto": "presupuesto|tabla|hoja|excel", "control de gastos": "tabla|hoja|excel",
    "código": "codigo|script|programa", "script": "script|codigo|programa",
    "programa": "programa|codigo|script", "página web": "pagina|web",
    "currículum": "curriculum|curriculo|cv", "mensaje": "mensaje|correo|email",
}
# Las órdenes que se entienden sin nombrar el documento, si es de hace un rato
# y fue lo último que dijo JARVIS. «Ponlo» o «pásalo» a secas NO: pueden ser
# cualquier cosa.
_RX_CLITICO_SEG = re.compile(
    r"\b(?:amplialo|ampliala|alargalo|alargala|acortalo|acortala|resumelo|resumela"
    r"|corrigelo|corrigela|mejoralo|mejorala|reescribelo|reescribela|traducelo|traducela"
    r"|abrelo|abrela|abremelo|abremela|muestramelo|muestramela|mandamelo|mandamela"
    r"|enviamelo|enviamela|conviertelo|conviertela|exportalo|exportala|anadele|agregale"
    r"|quitale|cambiale|incluyele|leemelo|leemela|completalo|completala|continualo"
    r"|continuala|terminalo|terminala|editalo|editala|modificalo|modificala|actualizalo"
    r"|actualizala|arreglalo|arreglala"
    r"|haz(?:lo|la)\s+(?:mas|menos|en|con|un|una|que|para|de|tipo|estilo)"
    r"|pas(?:a|ame)(?:lo|la)\s+(?:a|en)\s+(?:pdf|word|excel|txt|power ?point|markdown|html|rtf|csv"
    r"|(?:bloc|block|blog) de notas))\b")
_RX_SEG_ABRIR = re.compile(r"^\W*(?:vuelve a )?(?:abre|abreme|muestra|muestrame|ensena"
                           r"|ensename)(?:lo|la|melo|mela)?\b")
_RX_SEG_DONDE = re.compile(r"\bdonde (?:esta|quedo|se guardo|lo guardaste|la guardaste"
                           r"|lo dejaste|la dejaste|lo pusiste|la pusiste|lo tengo|la tengo)\b")
_RX_SEG_ENVIAR = re.compile(r"\b(?:manda|envia|pasa)(?:me)?(?:lo|la)?\b.*\b(?:movil|telefono"
                            r"|celular|cel)\b")
_RX_SEG_LEER = re.compile(r"^\W*(?:lee|leeme|leemelo|leemela|leelo|leela)\b")
_RX_SEG_FORMATO = re.compile(
    r"\b(?:en|a|como)\s+(?:un\s+|una\s+|formato\s+)?(?P<f>pdf|word|docx|excel|xlsx|txt"
    r"|texto plano|(?:bloc|block|blog) de notas|markdown|power ?point|pptx|presentacion"
    r"|html|rtf|csv)\b")
_RX_SEG_CAMBIO = re.compile(
    r"\b(?:anade|anadele|agrega|agregale|incluye|incluyele|pon|ponle|mete|metele|quita"
    r"|quitale|elimina|borra|cambia|cambiale|corrige|corrigelo|corrigela|mejora|mejoralo"
    r"|mejorala|amplia|amplialo|ampliala|alarga|alargalo|alargala|extiende|acorta"
    r"|acortalo|acortala|resume|resumelo|resumela|reescribe|reescribelo|reescribela"
    r"|traduce|traducelo|traducela|haz|hazlo|hazla|ajusta|modifica|modificalo|edita"
    r"|editalo|actualiza|actualizalo|completa|completalo|continua|continualo|sigue"
    r"|siguelo|termina|terminalo|arregla|arreglalo|arreglala|sustituye|reemplaza)\b")


def _rx_referencia(ult: dict):
    """«el documento», «la carta», «ese poema»… referidos al ÚLTIMO documento.

    Solo el tipo que era de verdad y los nombres genéricos, y nunca seguidos de
    algo que señale a otro archivo: «lee el correo» no es «lee el documento», y
    «abre el archivo de configuración» no es el último documento.
    """
    tipo = _plano(ult.get("tipo") or "")
    propios = _SINONIMOS.get(ult.get("tipo") or "", re.escape(tipo) if tipo else "")
    nombres = _GENERICOS + (f"|{propios}" if propios else "")
    return re.compile(
        r"\b(?:el|la|al|del|ese|esa|este|esta|eso|lo del)\s+(?:mismo\s+|misma\s+|ultimo\s+"
        r"|ultima\s+)?(?:" + nombres + r")\b(?!\s+(?:de|del|llamad[oa]|que se llama|sobre"
        r"|acerca|nuevo|nueva)\b)(?!\s+\S+\.\w{2,4}\b)"
        # «el texto que tengo copiado», «el que te pasé»: otro texto, no el suyo.
        r"(?!\s+que\s+(?:tengo|te\s+\w+|hay|esta|aparece|sale|ves|me\s+mandaron))")


def _entender_seguimiento(orig: str, p: str, core):
    """«Ábrelo», «hazlo más largo», «pásalo a PDF» sobre el último documento."""
    ult = ultimo()
    if not ult or not ult.get("ruta") or not os.path.isfile(ult["ruta"]):
        return None
    edad = time.time() - float(ult.get("ts") or 0)
    # Nombrándolo («la carta»), vale durante el día; sin nombrarlo («ábrelo»),
    # solo si acaba de entregarse y no se ha hablado de otra cosa después.
    ref = edad < 12 * 3600 and _rx_referencia(ult).search(p)
    reciente = edad < _RECIENTE_S and _fue_lo_ultimo(core, ult)
    if not ref and not (reciente and _RX_CLITICO_SEG.search(p)):
        return None
    # «hazme OTRO documento sobre…» es uno nuevo, no un cambio de este.
    if re.search(r"\b(?:otro|otra|nuevo|nueva)\s+(?:" + _GENERICOS + r"|carta|poema|ensayo"
                 r"|informe|presentacion|tabla|lista|cuento|script|codigo)\b", p):
        return None
    if _RX_MENSAJERIA.search(p) and not _RX_SEG_ENVIAR.search(p):
        return None
    if re.search(r"\b(?:portapapeles|pantalla|camara|foto|imagen|captura)\b", p):
        return None
    base = {"modo": "seguimiento", "frase": orig, "ultimo": ult}
    if _RX_SEG_DONDE.search(p):
        return dict(base, accion="donde")
    if _RX_SEG_ENVIAR.search(p):
        return dict(base, accion="enviar")
    m_f = _RX_SEG_FORMATO.search(p)
    if m_f and re.search(r"\b(?:guarda|pasa|convierte|exporta|pon|haz|dame|quiero"
                         r"|necesito|copia)\w*\b", p):
        f = _plano(m_f.group("f"))
        nuevo = {"word": "docx", "excel": "xlsx", "texto plano": "txt", "presentacion": "pptx",
                 "powerpoint": "pptx", "power point": "pptx"}.get(f, f)
        if re.match(r"(?:bloc|block|blog) de notas?", f):
            nuevo = "txt"
        if nuevo != ult.get("formato"):
            return dict(base, accion="convertir", formato=nuevo)
    if _RX_SEG_ABRIR.search(p) and len(p.split()) <= 9:
        return dict(base, accion="abrir")
    if _RX_SEG_LEER.search(p) and len(p.split()) <= 8:
        return dict(base, accion="leer")
    if _RX_SEG_CAMBIO.search(p):
        # «borra el documento» entero no es un cambio: eso es «deshaz eso».
        if re.search(r"\b(?:borra|elimina)(?:lo|la)?\s*(?:el|la|ese|esa|este|esta)?\s*"
                     r"(?:" + _GENERICOS + r")?\s*$", p):
            return None
        return dict(base, accion="modificar")
    return None


# ══════════════════════════════════════════════════════════════════════════
# EL CEREBRO: la cadena del núcleo, con relevo
# ══════════════════════════════════════════════════════════════════════════
_ULTIMO_FALLO = [""]

_PIDE_CONTINUAR = ("Te has quedado cortado. Continúa exactamente desde la última "
                   "palabra que escribiste, sin repetir nada, sin saludar y sin "
                   "anunciar que continúas.")


def _es_local(url: str) -> bool:
    try:
        import cerebro_local
        return bool(cerebro_local.es_local(url))
    except Exception:
        u = (url or "").lower()
        return any(h in u for h in ("localhost", "127.0.0.1", "0.0.0.0", "[::1]"))


def _cliente(core, url: str, clave: str):
    hecho = getattr(core, "_cliente_llm", None)
    if callable(hecho):
        try:
            return hecho(url, clave)
        except Exception:
            pass
    try:
        from proveedor_claude import cliente as OpenAI
    except Exception:
        from openai import OpenAI
    return OpenAI(base_url=url, api_key=clave)


def _sin_pensamiento(texto: str) -> str:
    try:
        from proveedor_claude import sin_pensamiento
        texto = sin_pensamiento(texto)
    except Exception:
        pass
    texto = re.sub(r"<think>.*?</think>", "", texto or "", flags=re.S)
    if "</think>" in texto:
        texto = texto.split("</think>", 1)[1]
    if "<think>" in texto:
        texto = texto.split("<think>", 1)[0]
    return texto


def _llamar(cliente, modelo: str, mensajes: list, tope: int, temperatura: float):
    """Una llamada, adaptándose a lo que acepte cada proveedor."""
    comun = {"model": modelo, "messages": mensajes, "max_tokens": tope,
             "temperature": temperatura}
    extra = {"extra_body": {"reasoning_effort": "low"}, "timeout": _TIMEOUT}
    respuesta = None
    for _ in range(6):
        try:
            respuesta = cliente.chat.completions.create(**comun, **extra)
            break
        except TypeError:
            # SDK o traductor sin extra_body/timeout
            if not extra:
                raise
            extra = {}
        except Exception as e:
            if type(e).__name__ in ("APITimeoutError", "APIConnectionError", "ConnectError",
                                    "ConnectTimeout", "ReadTimeout"):
                raise
            err = str(e).lower()
            if comun["max_tokens"] > 4096 and re.search(
                    r"max_?tokens|max(?:imum)?[ _](?:output|completion|new)|too large"
                    r"|context.?length", err):
                comun["max_tokens"] = 4096
                continue
            if "temperature" in err and "temperature" in comun:
                comun.pop("temperature")
                continue
            if "extra_body" in extra:
                extra.pop("extra_body")
                continue
            raise
    if respuesta is None:
        raise RuntimeError("el proveedor no aceptó la petición")
    eleccion = respuesta.choices[0]
    texto = _sin_pensamiento(getattr(eleccion.message, "content", "") or "").strip()
    return texto, (getattr(eleccion, "finish_reason", "") or ""), respuesta


def _anotar_uso(nombre: str, url: str, modelo: str, respuesta, log):
    try:
        import metricas
        metricas.anotar_cerebro(url, modelo, respuesta)
    except Exception:
        pass
    try:
        uso = getattr(respuesta, "usage", None)
        if uso:
            import presupuesto
            presupuesto.registrar_uso(nombre, modelo, getattr(uso, "prompt_tokens", 0) or 0,
                                      getattr(uso, "completion_tokens", 0) or 0, log=log)
    except Exception:
        pass


def _empalmar(texto: str, mas: str) -> str:
    """Pega la continuación sin repetir el solape que a veces mete el modelo."""
    mas = (mas or "").lstrip()
    for n in range(min(len(texto), len(mas), 200), 15, -1):
        if texto.endswith(mas[:n]):
            return texto + mas[n:]
    separador = "" if texto.endswith((" ", "\n")) or mas[:1] in ".,;:!?)»" else " "
    return texto + separador + mas


def preguntar(core, sistema: str, usuario: str, tope: int = 3000,
              temperatura: float = 0.6, continuar: int = 2, log=print) -> str:
    """El texto del primer cerebro que conteste. '' si ninguno pudo.

    Recorre `core._proveedores()` en orden, respetando el tope de gasto y la
    cuarentena de cerebro_salud, igual que el chat: si Anthropic no tiene saldo
    o Pollinations está caído, escribe el Qwen de casa. Si el modelo se queda
    sin tokens a mitad de documento, se le pide que siga.
    """
    mensajes = [{"role": "system", "content": sistema},
                {"role": "user", "content": usuario}]
    try:
        proveedores = list(core._proveedores()) if core is not None else []
    except Exception as e:
        log(f"[REDACTOR] No pude leer los proveedores: {e}")
        proveedores = []

    if not proveedores:
        # Sin núcleo (el generador de la web puede ir suelto): el camino antiguo.
        try:
            import proveedor_claude
            texto = (proveedor_claude.responder(usuario, sistema=sistema, max_tokens=tope,
                                                timeout=_TIMEOUT, log=log) or "").strip()
            if texto:
                return _sin_pensamiento(texto).strip()
        except Exception as e:
            _ULTIMO_FALLO[0] = f"no tengo ningún cerebro configurado ({e})"
            return ""
        _ULTIMO_FALLO[0] = "no tengo ningún cerebro configurado"
        return ""

    ultimo_error = ""
    for nombre, url, modelo, clave in proveedores:
        if not _es_local(url):
            try:
                import presupuesto
                if not presupuesto.permite_nube():
                    log(f"[REDACTOR] Salto «{nombre}»: tope de gasto del día alcanzado")
                    ultimo_error = ultimo_error or "el tope de gasto del día está alcanzado"
                    continue
            except Exception:
                pass
        try:
            import cerebro_salud
            if cerebro_salud.en_cuarentena(nombre):
                log(f"[REDACTOR] Salto «{nombre}»: en cuarentena")
                continue
        except Exception:
            pass
        try:
            cli = _cliente(core, url, clave)
            log(f"[REDACTOR] Escribiendo con {nombre} ({modelo}), hasta {tope} tokens")
            texto, motivo, resp = _llamar(cli, modelo, mensajes, tope, temperatura)
            _anotar_uso(nombre, url, modelo, resp, log)
            if not texto and motivo == "length":
                # El razonamiento se comió el turno entero.
                texto, motivo, resp = _llamar(cli, modelo, mensajes, min(tope * 2, 32000),
                                              temperatura)
                _anotar_uso(nombre, url, modelo, resp, log)
            vueltas = 0
            while texto and motivo == "length" and vueltas < continuar:
                vueltas += 1
                log(f"[REDACTOR] Cortado por longitud; le pido que siga ({vueltas})")
                seguir = mensajes + [{"role": "assistant", "content": texto},
                                     {"role": "user", "content": _PIDE_CONTINUAR}]
                mas, motivo, resp = _llamar(cli, modelo, seguir, tope, temperatura)
                _anotar_uso(nombre, url, modelo, resp, log)
                if not mas:
                    break
                texto = _empalmar(texto, mas)
            if texto:
                try:
                    import cerebro_salud
                    cerebro_salud.registrar_exito(nombre)
                except Exception:
                    pass
                _ULTIMO_FALLO[0] = ""
                return texto
            ultimo_error = f"{nombre} devolvió una respuesta vacía"
            log(f"[REDACTOR] {ultimo_error}")
        except Exception as e:
            ultimo_error = f"{nombre}: {str(e)[:160]}"
            log(f"[REDACTOR] El proveedor «{nombre}» falló: {e}")
            try:
                import cerebro_salud
                cerebro_salud.registrar_fallo(nombre, str(e), log=log)
            except Exception:
                pass
    _ULTIMO_FALLO[0] = ultimo_error or "ningún cerebro contestó"
    return ""


def ultimo_fallo() -> str:
    """Por qué falló la última redacción, dicho para el señor."""
    crudo = _ULTIMO_FALLO[0]
    if not crudo:
        return ""
    try:
        from arrancar_ambos import diagnostico_cerebro
        que, como = diagnostico_cerebro(crudo)
        if que:
            return f"{que}. Para arreglarlo, {como}"
    except Exception:
        pass
    return crudo


# ══════════════════════════════════════════════════════════════════════════
# QUÉ SE LE PIDE AL CEREBRO
# ══════════════════════════════════════════════════════════════════════════
_SISTEMA_TEXTO = """Eres JARVIS y ahora haces de redactor profesional para el señor.
Escribe el texto COMPLETO que te encarga, listo para usarlo tal cual.

Reglas que no se saltan:
- En español, salvo que el encargo pida otro idioma.
- Entrega SOLO el texto final: nada de «Aquí tiene», ni comentarios, ni preguntas al final, ni explicaciones de lo que has hecho.
- La primera línea es el título, precedido de «# ». Usa «## » para los apartados cuando el tipo de texto los tenga (un poema, una carta o un mensaje no los llevan).
- Párrafos separados por una línea en blanco. En un poema o una canción, cada verso en su línea y las estrofas separadas por una línea en blanco.
- Para enumerar, «- » o «1. ». Negrita con **así**, con moderación. Si hace falta una tabla, en formato Markdown con barras verticales. Nada de HTML.
- Termina el texto: no lo dejes a medias ni pongas «[continuará]».
- Si falta un dato que solo sabe el señor (su nombre, una fecha, una dirección, un importe), déjalo entre corchetes, por ejemplo [Su nombre]. No lo inventes."""

_SISTEMA_CODIGO = """Eres un programador experto. Escribe el archivo COMPLETO y funcional que pide el señor, en {lenguaje}.
Reglas:
- Responde SOLO con el contenido del archivo: sin explicaciones fuera del código y sin bloques ```.
- Código limpio, que funcione a la primera, con comentarios breves en español.
- Nada de «TODO» ni de funciones vacías: impleméntalo entero.
- Si necesita librerías externas, indícalo en un comentario al principio con el comando para instalarlas.{extra}"""

_EXTRA_HTML = """
- Es una página web completa en UN solo archivo: <!DOCTYPE html>, el CSS dentro de <style> y el JavaScript dentro de <script>.
- Diseño moderno, legible y adaptable al móvil. Los textos de la página, en español."""

_SISTEMA_TABLA = """Eres JARVIS preparando una hoja de cálculo para el señor.
Responde SOLO con JSON válido, sin nada antes ni después:
{"titulo": "...", "hojas": [{"nombre": "...", "columnas": ["...", "..."], "filas": [[...], [...]], "totales": false}]}
Reglas:
- Datos útiles y realistas para el encargo. Si son datos que solo sabe el señor, deja la estructura lista con filas de ejemplo claramente marcadas como ejemplo.
- Los números van como números (sin símbolo de moneda ni separador de miles). Las fechas, AAAA-MM-DD.
- Si una columna se calcula con otras, pon la fórmula de Excel como texto que empiece por «=»: la fila 1 son los encabezados y los datos empiezan en la fila 2, así que la fila n usa sus celdas (p. ej. =B2*C2 en la fila 2).
- "totales": true si tiene sentido sumar al final las columnas numéricas.
- Entre 5 y 40 filas salvo que el encargo pida otra cosa. Nombres de hoja de 31 caracteres como máximo."""

_SISTEMA_DIAPOS = """Eres JARVIS preparando una presentación para el señor.
Responde SOLO con JSON válido, sin nada antes ni después:
{"titulo": "...", "subtitulo": "...", "diapositivas": [{"titulo": "...", "puntos": ["...", "..."], "notas": "..."}]}
Reglas:
- Entre 6 y 12 diapositivas salvo que el encargo pida otra cifra; la última, de conclusiones o cierre.
- De 3 a 6 puntos por diapositiva, frases cortas (15 palabras como mucho) con contenido real, no títulos vacíos.
- «notas» es lo que diría el orador en esa diapositiva: de 2 a 4 frases.
- En español salvo que se pida otro idioma."""

_SISTEMA_CAMBIO = """Eres JARVIS y estás corrigiendo un documento que ya escribiste para el señor.
Aplica EXACTAMENTE el cambio que pide y devuelve el documento ENTERO ya cambiado, con el mismo formato que tenía (título con «# », apartados con «## »…). Si el documento era JSON, devuelve el JSON entero. Si era código, el archivo entero.
Solo el documento: sin comentarios antes ni después."""


def _guion_tipo(tipo: str) -> str:
    """Una línea de oficio según el género: lo que distingue un poema de un acta."""
    return {
        "poema": "Un poema de verdad: imágenes concretas, ritmo y, si encaja, rima. Nada de ripios.",
        "canción": "Letra de canción con estrofas y estribillo marcado.",
        "carta": "Carta con saludo, cuerpo y despedida; tono adecuado al destinatario.",
        "carta formal": "Carta formal: lugar y fecha, destinatario, asunto, saludo, cuerpo, despedida y firma.",
        "ensayo": "Ensayo con introducción, desarrollo argumentado en apartados y conclusión.",
        "informe": "Informe con resumen, apartados claros, datos y conclusiones o recomendaciones.",
        "reporte": "Reporte con resumen, hallazgos y conclusiones.",
        "currículum": "Currículum con datos de contacto, perfil, experiencia, formación y habilidades; los datos personales entre corchetes.",
        "contrato": "Contrato con partes, antecedentes, cláusulas numeradas y firmas; los datos de las partes entre corchetes.",
        "cuento": "Relato con planteamiento, nudo y desenlace, y personajes con nombre.",
        "guion": "Guion con escenas, acotaciones y diálogos con el nombre de quien habla.",
        "discurso": "Discurso para decirlo en voz alta: frases claras, cercano y con un cierre que se recuerde.",
        "lista": "Lista concreta y útil; cada punto con una línea de explicación si aporta.",
        "resumen": "Resumen fiel y bien ordenado de lo esencial.",
        "receta": "Ingredientes con cantidades y pasos numerados.",
        "plan": "Plan con objetivos, pasos o fases, tiempos y recursos.",
        "cuestionario": "Preguntas numeradas y, al final, las respuestas correctas.",
        "acta": "Acta con fecha, asistentes, orden del día, acuerdos y firmas.",
    }.get(tipo, "")


def _encargo_texto(enc: dict, texto_base: str = "") -> tuple:
    usuario = [f"Encargo del señor: «{enc['encargo']}»",
               f"Tipo de texto: {enc.get('tipo') or 'documento'}"]
    if enc.get("palabras"):
        usuario.append(f"Extensión aproximada: {enc['palabras']} palabras.")
    guion = _guion_tipo(enc.get("tipo", ""))
    if guion:
        usuario.append(guion)
    if texto_base:
        usuario.append("Texto de partida (lo que JARVIS le acaba de contestar al señor):\n<<<\n"
                       + texto_base[:20000] + "\n>>>")
    tope = max(1200, min(int((enc.get("palabras") or 600) * 2.2) + 900, 16000))
    return "\n".join(usuario), tope


def _quitar_preambulo(texto: str) -> str:
    """El «¡Claro! Aquí tiene su poema:» que se cuela aunque se le diga."""
    partes = re.split(r"\n\s*\n", (texto or "").strip())
    if len(partes) > 1:
        primero = _plano(partes[0]).strip()
        if len(primero) < 220 and re.match(
                r"^(?:¡?claro|por supuesto|desde luego|con (?:mucho )?gusto|aqui (?:tiene|tienes"
                r"|esta|va|le dejo)|a continuacion|enseguida|perfecto|muy bien|de acuerdo"
                r"|vale|listo|senor)\b", primero):
            partes = partes[1:]
    if len(partes) > 1:
        ultimo_p = _plano(partes[-1]).strip()
        if len(ultimo_p) < 260 and re.match(
                r"^(?:¿|espero|si (?:quiere|desea|necesita|lo desea)|quiere que|desea que"
                r"|digame|hagame saber|no dude)", ultimo_p):
            partes = partes[:-1]
    texto = "\n\n".join(partes).strip()
    # Un bloque ```markdown``` que envuelve TODO el texto
    m = re.match(r"^```(?:markdown|md|text|texto)?\s*\n(.*)\n```\s*$", texto, re.S)
    return m.group(1).strip() if m else texto


def _limpiar_codigo(texto: str) -> str:
    bloques = re.findall(r"```[\w+#.-]*[ \t]*\n(.*?)```", texto or "", re.S)
    if bloques:
        return max(bloques, key=len).strip("\n") + "\n"
    return (texto or "").strip("\n") + "\n"


def _rescatar_json(texto: str):
    try:
        import pensar
        return pensar._rescatar_json(texto)
    except Exception:
        pass
    if not texto:
        return None
    m = re.search(r"[{\[].*[}\]]", texto, re.S)
    if not m:
        return None
    try:
        return json.loads(m.group(0))
    except Exception:
        return None


def _revisar_codigo(ext: str, codigo: str) -> str:
    """'' si compila o no hay forma barata de saberlo. No ejecuta nada."""
    try:
        if ext == "py":
            compile(codigo, "<jarvis>", "exec")
        elif ext == "json":
            json.loads(codigo)
    except SyntaxError as e:
        return f"línea {e.lineno}: {e.msg}"
    except Exception as e:
        return str(e)[:160]
    return ""


def _redactar(core, enc: dict, log=print, texto_base: str = ""):
    """El contenido según la familia. None si el cerebro no dio nada útil."""
    fam = familia(enc["formato"])
    if fam == "texto":
        usuario, tope = _encargo_texto(enc, texto_base)
        crudo = preguntar(core, _SISTEMA_TEXTO, usuario, tope=tope, log=log)
        texto = _quitar_preambulo(crudo)
        return texto if len(texto.strip()) >= 10 else None
    if fam == "codigo":
        ext = enc["formato"]
        sistema = _SISTEMA_CODIGO.format(lenguaje=_LENGUAJES.get(ext, ext),
                                         extra=_EXTRA_HTML if ext == "html" else "")
        usuario = f"Encargo del señor: «{enc['encargo']}»"
        if texto_base:
            usuario += f"\n\nMaterial de partida:\n<<<\n{texto_base[:20000]}\n>>>"
        codigo = _limpiar_codigo(preguntar(core, sistema, usuario, tope=8000,
                                           temperatura=0.2, log=log))
        if len(codigo.strip()) < 5:
            return None
        error = _revisar_codigo(ext, codigo)
        if error:
            log(f"[REDACTOR] El código no compila ({error}); le pido que lo corrija")
            arreglo = _limpiar_codigo(preguntar(
                core, sistema, usuario + "\n\nEsta versión tiene un error («" + error
                + "»). Devuélvela entera y corregida:\n<<<\n" + codigo + "\n>>>",
                tope=8000, temperatura=0.1, log=log))
            if len(arreglo.strip()) > 5 and not _revisar_codigo(ext, arreglo):
                codigo = arreglo
        return codigo
    sistema = _SISTEMA_TABLA if fam == "tabla" else _SISTEMA_DIAPOS
    usuario = f"Encargo del señor: «{enc['encargo']}»"
    if texto_base:
        usuario += f"\n\nMaterial de partida:\n<<<\n{texto_base[:20000]}\n>>>"
    datos = _rescatar_json(preguntar(core, sistema, usuario, tope=6000, temperatura=0.3,
                                     log=log))
    if datos is None:
        datos = _rescatar_json(preguntar(core, sistema + "\n\nResponde SOLO con el JSON.",
                                         usuario, tope=6000, temperatura=0.0, log=log))
    if isinstance(datos, list):
        datos = {"hojas": [{"filas": datos}]} if fam == "tabla" else {"diapositivas": datos}
    if not isinstance(datos, dict):
        return None
    if fam == "tabla" and not (datos.get("hojas") or datos.get("filas")):
        return None
    if fam == "diapos" and not datos.get("diapositivas"):
        return None
    return datos


# ══════════════════════════════════════════════════════════════════════════
# ESCRIBIR: del Markdown sencillo a cada formato
# ══════════════════════════════════════════════════════════════════════════
def _bloques(md: str) -> list:
    """[{t: titulo|h2|h3|p|li|ol|cita|codigo|hr|tabla, x: texto, ...}]

    Un párrafo conserva sus saltos de línea: en un poema, cada verso es una
    línea y juntarlos todos en una sola sería destrozarlo.
    """
    salida, parrafo, codigo, tabla = [], [], [], []
    en_codigo = False

    def cerrar():
        if parrafo:
            salida.append({"t": "p", "x": "\n".join(parrafo)})
            parrafo.clear()
        if tabla:
            filas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in tabla
                     if not re.match(r"^\|?[\s:|-]+\|?$", f.strip())]
            if filas:
                salida.append({"t": "tabla", "filas": filas})
            tabla.clear()

    for linea in (md or "").replace("\r\n", "\n").split("\n"):
        s = linea.strip()
        if s.startswith("```"):
            if en_codigo:
                salida.append({"t": "codigo", "x": "\n".join(codigo)})
                codigo, en_codigo = [], False
            else:
                cerrar()
                en_codigo = True
            continue
        if en_codigo:
            codigo.append(linea.rstrip())
            continue
        if not s:
            cerrar()
            continue
        if s.startswith("|") and s.count("|") >= 2:
            if parrafo:
                salida.append({"t": "p", "x": "\n".join(parrafo)})
                parrafo.clear()
            tabla.append(s)
            continue
        if tabla:
            cerrar()
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            cerrar()
            n = len(m.group(1))
            salida.append({"t": "titulo" if n == 1 else "h2" if n == 2 else "h3",
                           "x": m.group(2).strip().strip("#").strip()})
            continue
        if re.match(r"^(?:-{3,}|\*{3,}|_{3,})$", s):
            cerrar()
            salida.append({"t": "hr", "x": ""})
            continue
        m = re.match(r"^[-*•·]\s+(.*)$", s)
        if m:
            cerrar()
            salida.append({"t": "li", "x": m.group(1).strip()})
            continue
        m = re.match(r"^(\d{1,3})[.)]\s+(.*)$", s)
        if m:
            cerrar()
            salida.append({"t": "ol", "x": m.group(2).strip(), "n": int(m.group(1))})
            continue
        m = re.match(r"^>\s?(.*)$", s)
        if m:
            cerrar()
            salida.append({"t": "cita", "x": m.group(1).strip()})
            continue
        parrafo.append(s)
    if en_codigo and codigo:
        salida.append({"t": "codigo", "x": "\n".join(codigo)})
    cerrar()
    # Un título en negrita suelta en la primera línea cuenta como título.
    if salida and salida[0]["t"] == "p":
        m = re.fullmatch(r"\*\*(.{2,120}?)\*\*", salida[0]["x"].strip())
        if m:
            salida[0] = {"t": "titulo", "x": m.group(1).strip()}
    return salida


_RX_INLINE = re.compile(r"(\*\*[^*\n]+?\*\*|`[^`\n]+?`|(?<![\w*])\*[^*\n]+?\*(?![\w*]))")


def _trozos(texto: str) -> list:
    """[(texto, negrita, cursiva, codigo)] de una línea con Markdown sencillo."""
    salida = []
    for parte in _RX_INLINE.split(texto or ""):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            salida.append((parte[2:-2], True, False, False))
        elif parte.startswith("`") and parte.endswith("`") and len(parte) > 2:
            salida.append((parte[1:-1], False, False, True))
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            salida.append((parte[1:-1], False, True, False))
        else:
            salida.append((parte, False, False, False))
    return salida


def _sin_marcas(texto: str) -> str:
    return "".join(t for t, _n, _c, _k in _trozos(texto))


def titulo_de(md: str) -> str:
    for b in _bloques(md):
        if b["t"] in ("titulo", "h2"):
            return _sin_marcas(b["x"]).strip()[:120]
        break
    return ""


def _texto_tabla(filas: list) -> list:
    anchos = [0] * max(len(f) for f in filas)
    for f in filas:
        for i, c in enumerate(f):
            anchos[i] = min(max(anchos[i], len(_sin_marcas(c))), 40)
    lineas = []
    for k, f in enumerate(filas):
        celdas = [_sin_marcas(c).ljust(anchos[i]) for i, c in enumerate(f)]
        lineas.append("  ".join(celdas).rstrip())
        if k == 0:
            lineas.append("  ".join("-" * a for a in anchos))
    return lineas


def a_texto_plano(md: str) -> str:
    """El Markdown sencillo, pasado a texto que se lee bien en el Bloc de notas."""
    lineas = []
    previo = ""
    for b in _bloques(md):
        t = b["t"]
        x = _sin_marcas(b.get("x", ""))
        if previo in ("li", "ol") and t not in ("li", "ol"):
            lineas.append("")
        if t == "titulo":
            lineas += [x, "=" * min(max(len(x), 3), 72), ""]
        elif t == "h2":
            lineas += ["", x, "-" * min(max(len(x), 3), 72), ""]
        elif t == "h3":
            lineas += ["", x, ""]
        elif t == "p":
            lineas += [x, ""]
        elif t == "li":
            lineas.append(f"  • {x}")
        elif t == "ol":
            lineas.append(f"  {b.get('n', 1)}. {x}")
        elif t == "cita":
            lineas += [f"    {x}" if x[:1] in "«\"“'" else f"    «{x}»", ""]
        elif t == "codigo":
            lineas += [b.get("x", ""), ""]
        elif t == "hr":
            lineas += ["", "* * *", ""]
        elif t == "tabla":
            lineas += _texto_tabla(b["filas"]) + [""]
        previo = t
    texto = "\n".join(lineas)
    return re.sub(r"\n{3,}", "\n\n", texto).strip() + "\n"


def _escribir_txt(ruta: str, md: str, literal: bool = False):
    texto = md if literal else a_texto_plano(md)
    # Con BOM: el Bloc de notas de antes de 2019 abre sin él en ANSI y
    # convierte cada tilde en un jeroglífico.
    with open(ruta, "w", encoding="utf-8-sig") as f:
        f.write(texto if texto.endswith("\n") else texto + "\n")


def _escribir_md(ruta: str, md: str):
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(md.strip() + "\n")


def _runs_docx(parrafo, texto: str, Pt=None):
    lineas = (texto or "").split("\n")
    for i, linea in enumerate(lineas):
        for trozo, negrita, cursiva, codigo in _trozos(linea):
            run = parrafo.add_run(trozo)
            run.bold = negrita or None
            run.italic = cursiva or None
            if codigo:
                run.font.name = "Consolas"
        if i < len(lineas) - 1:
            parrafo.add_run().add_break()


def _escribir_docx(ruta: str, md: str, titulo: str = ""):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    estilos = {s.name for s in doc.styles}
    bloques = _bloques(md)
    if titulo and not any(b["t"] == "titulo" for b in bloques):
        doc.add_heading(titulo, level=0)
    for b in bloques:
        t = b["t"]
        if t == "titulo":
            doc.add_heading(_sin_marcas(b["x"]), level=0)
        elif t == "h2":
            doc.add_heading(_sin_marcas(b["x"]), level=1)
        elif t == "h3":
            doc.add_heading(_sin_marcas(b["x"]), level=2)
        elif t == "p":
            _runs_docx(doc.add_paragraph(), b["x"])
        elif t == "cita":
            _runs_docx(doc.add_paragraph(style="Quote" if "Quote" in estilos else None), b["x"])
        elif t == "li":
            _runs_docx(doc.add_paragraph(style="List Bullet" if "List Bullet" in estilos else None),
                       b["x"])
        elif t == "ol":
            # El número escrito y no «List Number»: ese estilo sigue contando de
            # una lista a la siguiente y la segunda empezaría en 6.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-14)
            _runs_docx(p, f"{b.get('n', 1)}. " + b["x"])
        elif t == "codigo":
            p = doc.add_paragraph()
            run = p.add_run(b["x"])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif t == "hr":
            p = doc.add_paragraph("* * *")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif t == "tabla":
            filas = b["filas"]
            columnas = max(len(f) for f in filas)
            tabla = doc.add_table(rows=0, cols=columnas)
            if "Table Grid" in estilos:
                tabla.style = "Table Grid"
            for k, fila in enumerate(filas):
                celdas = tabla.add_row().cells
                for i in range(columnas):
                    texto = fila[i] if i < len(fila) else ""
                    celdas[i].text = ""
                    _runs_docx(celdas[i].paragraphs[0], f"**{_sin_marcas(texto)}**"
                               if k == 0 and texto else texto)
            doc.add_paragraph()
    doc.core_properties.title = titulo or titulo_de(md)
    doc.core_properties.author = "JARVIS"
    doc.save(ruta)


def _rtf(texto: str) -> str:
    salida = []
    for ch in texto:
        o = ord(ch)
        if ch in "\\{}":
            salida.append("\\" + ch)
        elif ch == "\n":
            salida.append("\\line ")
        elif ch == "\t":
            salida.append("\\tab ")
        elif o < 128:
            salida.append(ch)
        elif o < 0x10000:
            salida.append(f"\\u{o if o < 32768 else o - 65536}?")
        else:
            o -= 0x10000
            alto, bajo = 0xD800 + (o >> 10), 0xDC00 + (o & 0x3FF)
            salida.append(f"\\u{alto - 65536}?\\u{bajo - 65536}?")
    return "".join(salida)


def _rtf_inline(texto: str) -> str:
    partes = []
    for trozo, negrita, cursiva, codigo in _trozos(texto):
        cuerpo = _rtf(trozo)
        if codigo:
            cuerpo = "{\\f1 " + cuerpo + "}"
        if cursiva:
            cuerpo = "{\\i " + cuerpo + "}"
        if negrita:
            cuerpo = "{\\b " + cuerpo + "}"
        partes.append(cuerpo)
    return "".join(partes)


def _escribir_rtf(ruta: str, md: str, titulo: str = ""):
    """RTF a mano: lo abren Word, WordPad y LibreOffice sin librería ninguna."""
    cuerpo = ["{\\rtf1\\ansi\\ansicpg1252\\deff0{\\fonttbl{\\f0\\fswiss Calibri;}"
              "{\\f1\\fmodern Consolas;}}\\viewkind4\\uc1\\f0\\fs22"]
    bloques = _bloques(md)
    if titulo and not any(b["t"] == "titulo" for b in bloques):
        bloques.insert(0, {"t": "titulo", "x": titulo})
    for b in bloques:
        t, x = b["t"], b.get("x", "")
        if t == "titulo":
            cuerpo.append("{\\pard\\qc\\sa240\\b\\fs36 " + _rtf(_sin_marcas(x)) + "\\par}")
        elif t == "h2":
            cuerpo.append("{\\pard\\sb240\\sa120\\b\\fs28 " + _rtf(_sin_marcas(x)) + "\\par}")
        elif t == "h3":
            cuerpo.append("{\\pard\\sb180\\sa80\\b\\fs24 " + _rtf(_sin_marcas(x)) + "\\par}")
        elif t == "p":
            cuerpo.append("{\\pard\\sa160 " + _rtf_inline(x) + "\\par}")
        elif t == "cita":
            cuerpo.append("{\\pard\\li720\\sa160\\i " + _rtf_inline(x) + "\\par}")
        elif t == "li":
            cuerpo.append("{\\pard\\fi-280\\li560\\sa60 \\bullet\\tab " + _rtf_inline(x) + "\\par}")
        elif t == "ol":
            cuerpo.append("{\\pard\\fi-280\\li560\\sa60 " + str(b.get("n", 1)) + ".\\tab "
                          + _rtf_inline(x) + "\\par}")
        elif t == "codigo":
            cuerpo.append("{\\pard\\sa160\\f1\\fs19 " + _rtf(x) + "\\par}")
        elif t == "hr":
            cuerpo.append("{\\pard\\qc\\sa160 * * *\\par}")
        elif t == "tabla":
            for linea in _texto_tabla(b["filas"]):
                cuerpo.append("{\\pard\\f1\\fs19 " + _rtf(linea) + "\\par}")
            cuerpo.append("{\\pard\\par}")
    cuerpo.append("}")
    with open(ruta, "w", encoding="ascii", errors="replace") as f:
        f.write("\n".join(cuerpo))


def _fuente_pdf() -> tuple:
    """Una TTF con tildes y eñes de sobra; Helvetica si no hay ninguna."""
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.fonts import addMapping
        if "JarvisSans" in pdfmetrics.getRegisteredFontNames():
            return "JarvisSans", "JarvisSans-Bold", "JarvisSans-Italic"
        windir = os.environ.get("WINDIR", r"C:\Windows")
        candidatos = [
            (os.path.join(windir, "Fonts", "arial.ttf"), os.path.join(windir, "Fonts", "arialbd.ttf"),
             os.path.join(windir, "Fonts", "ariali.ttf")),
            (os.path.join(windir, "Fonts", "calibri.ttf"), os.path.join(windir, "Fonts", "calibrib.ttf"),
             os.path.join(windir, "Fonts", "calibrii.ttf")),
            ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
             "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
             "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
            ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf",
             "/Library/Fonts/Arial Italic.ttf"),
        ]
        for normal, negrita, cursiva in candidatos:
            if os.path.isfile(normal):
                pdfmetrics.registerFont(TTFont("JarvisSans", normal))
                pdfmetrics.registerFont(TTFont("JarvisSans-Bold",
                                               negrita if os.path.isfile(negrita) else normal))
                pdfmetrics.registerFont(TTFont("JarvisSans-Italic",
                                               cursiva if os.path.isfile(cursiva) else normal))
                addMapping("JarvisSans", 0, 0, "JarvisSans")
                addMapping("JarvisSans", 1, 0, "JarvisSans-Bold")
                addMapping("JarvisSans", 0, 1, "JarvisSans-Italic")
                addMapping("JarvisSans", 1, 1, "JarvisSans-Bold")
                return "JarvisSans", "JarvisSans-Bold", "JarvisSans-Italic"
    except Exception:
        pass
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"


def _pdf_inline(texto: str, ttf: bool) -> str:
    partes = []
    for trozo, negrita, cursiva, codigo in _trozos(texto):
        if not ttf:
            trozo = trozo.encode("cp1252", "replace").decode("cp1252")
        t = trozo.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        t = t.replace("\n", "<br/>")
        if codigo:
            t = f'<font face="Courier">{t}</font>'
        if cursiva:
            t = f"<i>{t}</i>"
        if negrita:
            t = f"<b>{t}</b>"
        partes.append(t)
    return "".join(partes)


def _escribir_pdf(ruta: str, md: str, titulo: str = ""):
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

    fuente, negrita, cursiva = _fuente_pdf()
    ttf = fuente == "JarvisSans"
    base = getSampleStyleSheet()
    cuerpo = ParagraphStyle("J-cuerpo", parent=base["BodyText"], fontName=fuente, fontSize=11,
                            leading=15.5, spaceAfter=7, alignment=TA_JUSTIFY)
    st_titulo = ParagraphStyle("J-titulo", parent=base["Title"], fontName=negrita, fontSize=20,
                               leading=24, spaceAfter=14, alignment=TA_CENTER,
                               textColor=colors.HexColor("#1F3A5F"))
    st_h2 = ParagraphStyle("J-h2", parent=base["Heading2"], fontName=negrita, fontSize=14,
                           leading=18, spaceBefore=10, spaceAfter=6,
                           textColor=colors.HexColor("#1F3A5F"))
    st_h3 = ParagraphStyle("J-h3", parent=base["Heading3"], fontName=negrita, fontSize=12,
                           leading=16, spaceBefore=8, spaceAfter=4)
    st_lista = ParagraphStyle("J-lista", parent=cuerpo, leftIndent=18, bulletIndent=6,
                              spaceAfter=3, alignment=0)
    st_cita = ParagraphStyle("J-cita", parent=cuerpo, fontName=cursiva, leftIndent=28,
                             rightIndent=20, textColor=colors.HexColor("#444444"))
    st_codigo = ParagraphStyle("J-codigo", parent=base["Code"], fontSize=8.5, leading=11)

    historia = []
    bloques = _bloques(md)
    if titulo and not any(b["t"] == "titulo" for b in bloques):
        bloques.insert(0, {"t": "titulo", "x": titulo})
    for b in bloques:
        t, x = b["t"], b.get("x", "")
        if t == "titulo":
            historia.append(Paragraph(_pdf_inline(_sin_marcas(x), ttf), st_titulo))
        elif t == "h2":
            historia.append(Paragraph(_pdf_inline(_sin_marcas(x), ttf), st_h2))
        elif t == "h3":
            historia.append(Paragraph(_pdf_inline(_sin_marcas(x), ttf), st_h3))
        elif t == "p":
            historia.append(Paragraph(_pdf_inline(x, ttf), cuerpo))
        elif t == "cita":
            historia.append(Paragraph(_pdf_inline(x, ttf), st_cita))
        elif t == "li":
            historia.append(Paragraph(_pdf_inline(x, ttf), st_lista, bulletText="•"))
        elif t == "ol":
            historia.append(Paragraph(_pdf_inline(x, ttf), st_lista,
                                      bulletText=f"{b.get('n', 1)}."))
        elif t == "codigo":
            historia.append(Preformatted(x, st_codigo))
        elif t == "hr":
            historia.append(Spacer(1, 10))
        elif t == "tabla":
            filas = b["filas"]
            n = max(len(f) for f in filas)
            datos = [[Paragraph(_pdf_inline(f[i] if i < len(f) else "", ttf), cuerpo)
                      for i in range(n)] for f in filas]
            historia.append(Spacer(1, 4))
            tabla = Table(datos, repeatRows=1)
            tabla.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            historia += [tabla, Spacer(1, 8)]
    if not historia:
        historia.append(Paragraph(" ", cuerpo))
    doc = SimpleDocTemplate(ruta, pagesize=A4, leftMargin=2.4 * cm, rightMargin=2.4 * cm,
                            topMargin=2.2 * cm, bottomMargin=2.2 * cm,
                            title=titulo or titulo_de(md) or "Documento", author="JARVIS")
    doc.build(historia)


_RX_NUMERO = re.compile(r"^-?\d+(?:\.\d+)?$")
_RX_FECHA = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_MONEDA = re.compile(r"precio|importe|total|cost[oe]|monto|saldo|sueldo|salario|pago|ingreso"
                     r"|gasto|valor|subtotal|iva|impuesto|\$|€|mxn|usd|eur", re.I)


def _celda(valor):
    if valor is None or isinstance(valor, (int, float, bool)):
        return valor
    s = str(valor).strip()
    if s.startswith("="):
        # Una fórmula con «|» es DDE: ejecutaría programas al abrir el libro.
        return "'" + s if "|" in s else s
    limpio = s.replace(" ", "")
    if _RX_NUMERO.match(limpio):
        try:
            return int(limpio) if re.match(r"^-?\d+$", limpio) and len(limpio) < 16 else float(limpio)
        except ValueError:
            return s
    if _RX_FECHA.match(s):
        try:
            import datetime
            return datetime.date.fromisoformat(s)
        except ValueError:
            return s
    return s


def _hojas(datos: dict) -> list:
    hojas = datos.get("hojas")
    if not hojas and (datos.get("filas") or datos.get("columnas")):
        hojas = [datos]
    salida = []
    for i, h in enumerate(hojas or []):
        if not isinstance(h, dict):
            continue
        columnas = [str(c) for c in (h.get("columnas") or [])]
        filas = []
        for f in h.get("filas") or []:
            if isinstance(f, dict):
                if not columnas:
                    columnas = [str(k) for k in f.keys()]
                f = [f.get(c) for c in columnas]
            elif not isinstance(f, (list, tuple)):
                f = [f]
            filas.append(list(f))
        if not columnas and filas:
            columnas = [f"Columna {j + 1}" for j in range(max(len(f) for f in filas))]
        if columnas:
            salida.append({"nombre": str(h.get("nombre") or f"Hoja {i + 1}"),
                           "columnas": columnas, "filas": filas,
                           "totales": bool(h.get("totales"))})
    return salida


def _escribir_xlsx(ruta: str, datos: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    libro.remove(libro.active)
    usados = set()
    cabecera = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", start_color="1F3A5F", end_color="1F3A5F")
    borde = Border(*(Side(style="thin", color="BBBBBB"),) * 4)
    for hoja in _hojas(datos) or [{"nombre": "Hoja 1", "columnas": ["Dato"], "filas": [],
                                   "totales": False}]:
        nombre = re.sub(r"[\[\]:*?/\\]", "", hoja["nombre"])[:31].strip() or "Hoja"
        base, n = nombre, 2
        while nombre.lower() in usados:
            nombre = f"{base[:28]} {n}"
            n += 1
        usados.add(nombre.lower())
        ws = libro.create_sheet(nombre)
        columnas = hoja["columnas"]
        ws.append(columnas)
        for fila in hoja["filas"]:
            ws.append([_celda(v) for v in fila[:len(columnas)]])
        ultima = ws.max_row
        n_cols = len(columnas)
        for c in range(1, n_cols + 1):
            celda = ws.cell(row=1, column=c)
            celda.font, celda.fill, celda.border = cabecera, relleno, borde
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        numericas = []
        for c in range(1, n_cols + 1):
            valores = [ws.cell(row=r, column=c).value for r in range(2, ultima + 1)]
            presentes = [v for v in valores if v not in (None, "")]
            es_num = bool(presentes) and all(
                isinstance(v, (int, float)) and not isinstance(v, bool)
                or (isinstance(v, str) and v.startswith("=")) for v in presentes)
            if es_num:
                numericas.append(c)
            moneda = es_num and _MONEDA.search(columnas[c - 1] or "")
            for r in range(2, ultima + 1):
                celda = ws.cell(row=r, column=c)
                celda.border = borde
                if moneda:
                    celda.number_format = "#,##0.00"
                elif hasattr(celda.value, "isoformat"):
                    celda.number_format = "yyyy-mm-dd"
        if hoja["totales"] and ultima >= 2 and numericas:
            fila_total = ultima + 1
            ws.cell(row=fila_total, column=1, value="Total").font = Font(bold=True)
            for c in numericas:
                if c == 1:
                    continue
                letra = get_column_letter(c)
                celda = ws.cell(row=fila_total, column=c, value=f"=SUM({letra}2:{letra}{ultima})")
                celda.font = Font(bold=True)
                celda.number_format = ws.cell(row=2, column=c).number_format
        for c in range(1, n_cols + 1):
            largo = max([len(str(ws.cell(row=r, column=c).value or ""))
                         for r in range(1, ws.max_row + 1)] or [8])
            ws.column_dimensions[get_column_letter(c)].width = max(10, min(largo + 3, 60))
        ws.freeze_panes = "A2"
        if ultima >= 2:
            ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{ultima}"
    libro.properties.title = str(datos.get("titulo") or "")
    libro.properties.creator = "JARVIS"
    libro.save(ruta)


def _escribir_csv(ruta: str, datos: dict):
    hojas = _hojas(datos)
    hoja = hojas[0] if hojas else {"columnas": [], "filas": []}
    # Con BOM: sin él Excel abre el CSV en ANSI y destroza las tildes.
    with open(ruta, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(hoja["columnas"])
        for fila in hoja["filas"]:
            w.writerow(["" if v is None else v for v in fila])


def _escribir_pptx(ruta: str, datos: dict):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    ancho_viejo = prs.slide_width
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    factor = prs.slide_width / ancho_viejo
    azul = RGBColor(0x1F, 0x3A, 0x5F)

    def _ensanchar(diapo):
        # La plantilla por defecto es 4:3; a 16:9 los marcos quedan arrimados
        # a la izquierda si no se estiran. Las CUATRO medidas a la vez: el marco
        # hereda la posición de la plantilla, y fijar solo la horizontal dejaba
        # la vertical y la altura a cero (título y viñetas amontonados arriba).
        for forma in diapo.placeholders:
            izq, arriba, ancho, alto = forma.left, forma.top, forma.width, forma.height
            forma.left, forma.top = int(izq * factor), arriba
            forma.width, forma.height = int(ancho * factor), alto

    def _acento(diapo, izq, arriba, ancho, alto):
        barra = diapo.shapes.add_shape(MSO_SHAPE.RECTANGLE, izq, arriba, ancho, alto)
        barra.fill.solid()
        barra.fill.fore_color.rgb = azul
        barra.line.fill.background()
        barra.shadow.inherit = False

    portada = prs.slides.add_slide(prs.slide_layouts[0])
    _ensanchar(portada)
    _acento(portada, 0, prs.slide_height - Inches(0.35), prs.slide_width, Inches(0.35))
    portada.shapes.title.text = str(datos.get("titulo") or "Presentación")
    for p in portada.shapes.title.text_frame.paragraphs:
        p.font.bold, p.font.color.rgb = True, azul
    if len(portada.placeholders) > 1:
        portada.placeholders[1].text = str(datos.get("subtitulo") or "")

    for d in (datos.get("diapositivas") or [])[:60]:
        if isinstance(d, str):
            d = {"titulo": d, "puntos": []}
        if not isinstance(d, dict):
            continue
        diapo = prs.slides.add_slide(prs.slide_layouts[1])
        _ensanchar(diapo)
        titulo = diapo.shapes.title
        titulo.text = str(d.get("titulo") or "")
        for p in titulo.text_frame.paragraphs:
            p.font.bold, p.font.color.rgb = True, azul
        _acento(diapo, diapo.placeholders[1].left, titulo.top + titulo.height,
                Inches(1.6), Inches(0.07))
        puntos = [str(x) for x in (d.get("puntos") or d.get("contenido") or []) if str(x).strip()]
        if isinstance(d.get("puntos"), str):
            puntos = [x.strip(" -•") for x in d["puntos"].split("\n") if x.strip()]
        marco = diapo.placeholders[1].text_frame
        marco.clear()
        tam = Pt(28 if len(puntos) <= 4 else 24 if len(puntos) <= 6 else 20)
        for k, punto in enumerate(puntos):
            parrafo = marco.paragraphs[0] if k == 0 else marco.add_paragraph()
            parrafo.text = _sin_marcas(punto)
            parrafo.font.size = tam
        notas = str(d.get("notas") or "").strip()
        if notas:
            diapo.notes_slide.notes_text_frame.text = notas
    prs.core_properties.title = str(datos.get("titulo") or "")
    prs.core_properties.author = "JARVIS"
    prs.save(ruta)


def _diapos_a_md(datos: dict) -> str:
    partes = [f"# {datos.get('titulo') or 'Presentación'}", ""]
    if datos.get("subtitulo"):
        partes += [str(datos["subtitulo"]), ""]
    for i, d in enumerate(datos.get("diapositivas") or [], 1):
        if isinstance(d, dict):
            partes += [f"## {i}. {d.get('titulo') or ''}", ""]
            partes += [f"- {p}" for p in (d.get("puntos") or [])]
            if d.get("notas"):
                partes += ["", f"> {d['notas']}"]
            partes.append("")
    return "\n".join(partes)


def _tabla_a_md(datos: dict) -> str:
    partes = [f"# {datos.get('titulo') or 'Tabla'}", ""]
    for h in _hojas(datos):
        partes += [f"## {h['nombre']}", "", "| " + " | ".join(h["columnas"]) + " |",
                   "|" + "---|" * len(h["columnas"])]
        partes += ["| " + " | ".join("" if v is None else str(v) for v in f) + " |"
                   for f in h["filas"]]
        partes.append("")
    return "\n".join(partes)


def escribir(ruta: str, formato: str, contenido, titulo: str = "", literal: bool = False,
             log=print) -> tuple:
    """Escribe el contenido en el formato pedido. Devuelve (ruta_final, aviso).

    Si falta la librería de un formato, no se queda sin documento: Word sin
    python-docx sale en RTF (que Word abre igual), Excel sin openpyxl en CSV,
    PDF sin reportlab en Word, y PowerPoint sin python-pptx en un guion de
    texto. El aviso dice qué pasó y cómo instalar lo que falta.
    """
    formato = (formato or "txt").lower().lstrip(".")
    os.makedirs(os.path.dirname(ruta) or ".", exist_ok=True)
    fam = familia(formato)
    aviso = ""

    def _cambiar(ext):
        return os.path.splitext(ruta)[0] + "." + ext

    if fam in ("tabla", "diapos") and not isinstance(contenido, dict):
        contenido = {"titulo": titulo, "hojas": [{"columnas": ["Texto"],
                                                 "filas": [[l] for l in str(contenido).splitlines()
                                                           if l.strip()]}]} \
            if fam == "tabla" else {"titulo": titulo, "diapositivas": [
                {"titulo": _sin_marcas(b["x"]), "puntos": []} for b in _bloques(str(contenido))
                if b["t"] in ("h2", "titulo")]}
    if formato == "xlsx":
        try:
            _escribir_xlsx(ruta, contenido)
            return ruta, aviso
        except ImportError:
            ruta = _cambiar("csv")
            aviso = "No tengo openpyxl, así que va en CSV (Excel lo abre). Instálelo con «pip install openpyxl»."
            formato = "csv"
    if formato == "csv":
        _escribir_csv(ruta, contenido)
        return ruta, aviso
    if formato == "pptx":
        try:
            _escribir_pptx(ruta, contenido)
            return ruta, aviso
        except ImportError:
            ruta = _cambiar("txt")
            _escribir_txt(ruta, _diapos_a_md(contenido))
            return ruta, ("No tengo python-pptx, así que le dejo el guion de las diapositivas en "
                          "texto. Instálelo con «pip install python-pptx».")
    if fam == "codigo":
        with open(ruta, "w", encoding="utf-8", newline="\r\n" if formato in ("bat", "ps1") else None) as f:
            f.write(str(contenido))
        return ruta, aviso
    md = str(contenido)
    if formato == "pdf":
        try:
            _escribir_pdf(ruta, md, titulo)
            return ruta, aviso
        except ImportError:
            ruta = _cambiar("docx")
            aviso = "No tengo reportlab para hacer el PDF; va en Word. Instálelo con «pip install reportlab»."
            formato = "docx"
    if formato == "docx":
        try:
            _escribir_docx(ruta, md, titulo)
            return ruta, aviso
        except ImportError:
            ruta = _cambiar("rtf")
            aviso = (aviso + " " if aviso else "") + ("No tengo python-docx, así que va en RTF "
                                                     "(Word lo abre igual). Instálelo con "
                                                     "«pip install python-docx».")
            formato = "rtf"
    if formato == "rtf":
        _escribir_rtf(ruta, md, titulo)
        return ruta, aviso.strip()
    if formato == "md":
        _escribir_md(ruta, md)
        return ruta, aviso
    _escribir_txt(ruta, md, literal=literal)
    return ruta, aviso


def leer_contenido(ruta: str, formato: str = "") -> str:
    """El texto de un documento ya escrito, para poder cambiarlo."""
    formato = (formato or os.path.splitext(ruta)[1]).lower().lstrip(".")
    try:
        if formato == "docx":
            from docx import Document
            partes = []
            for p in Document(ruta).paragraphs:
                estilo = (p.style.name or "").lower() if p.style is not None else ""
                texto = p.text.strip()
                if not texto:
                    continue
                if estilo == "title":
                    partes.append("# " + texto)
                elif estilo.startswith("heading 1"):
                    partes.append("## " + texto)
                elif estilo.startswith("heading"):
                    partes.append("### " + texto)
                elif "list bullet" in estilo:
                    partes.append("- " + texto)
                else:
                    partes.append(texto)
            return "\n\n".join(partes)
        if formato in ("pdf", "xlsx", "pptx", "rtf"):
            return ""
        with open(ruta, encoding="utf-8-sig", errors="replace") as f:
            return f.read()
    except Exception:
        return ""


# ══════════════════════════════════════════════════════════════════════════
# NOMBRES Y RUTAS
# ══════════════════════════════════════════════════════════════════════════
_RESERVADOS = {"con", "prn", "aux", "nul"} | {f"com{i}" for i in range(1, 10)} | \
    {f"lpt{i}" for i in range(1, 10)}


def nombre_archivo(base: str, formato: str) -> str:
    """Un nombre que Windows acepta, legible y con su extensión."""
    base = _sin_marcas(base or "").strip()
    base = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", base)
    base = re.sub(r"\s+", " ", base).strip(" .")
    if len(base) > 60:
        corte = base[:60].rsplit(" ", 1)[0]
        base = corte if len(corte) > 20 else base[:60]
    if not base:
        base = "Documento " + time.strftime("%Y-%m-%d %H%M")
    if base.lower() in _RESERVADOS:
        base += "_"
    formato = (formato or "txt").lower().lstrip(".")
    if not base.lower().endswith("." + formato):
        base += "." + formato
    return base


def nombre_codigo(descripcion: str, formato: str) -> str:
    """«script en Python que ordene archivos por fecha» -> ordene_archivos_por_fecha.py

    Un programa con espacios en el nombre no se puede importar ni llamar
    cómodamente desde la consola: para código, nombre en minúsculas y con
    guiones bajos.
    """
    p = _plano(descripcion or "")
    p = re.sub(r"^\s*(?:un|una|el|la)\s+", "", p)
    p = re.sub(r"^(?:script|codigo|programa|funcion|algoritmo|paginas? web|sitio web|pagina"
               r"|archivo)\b\s*", "", p)
    p = re.sub(r"^(?:en|de)\s+(?:" + "|".join(re.escape(_plano(n).split()[0].lower())
                                            for n in _LENGUAJES.values()) + r")\b\s*", "", p)
    p = re.sub(r"^(?:que|para|de|del|con|sobre)\s+", "", p)
    palabras = [w for w in re.findall(r"[a-z0-9]+", p)
                if w not in ("el", "la", "los", "las", "un", "una", "mi", "mis")][:6]
    base = "_".join(palabras) or ("pagina" if formato == "html" else "programa")
    return base[:60].strip("_") + "." + (formato or "txt")


def ruta_libre(carpeta_destino: str, nombre: str) -> str:
    """Nunca sobrescribe: «X.txt», «X (2).txt», «X (3).txt»…"""
    raiz, ext = os.path.splitext(nombre)
    candidato = os.path.join(carpeta_destino, nombre)
    n = 2
    while os.path.exists(candidato):
        candidato = os.path.join(carpeta_destino, f"{raiz} ({n}){ext}")
        n += 1
    return candidato


def ruta_documento(nombre: str) -> str:
    """Dónde poner un documento que llega con un nombre suelto («poema.txt»):
    en la carpeta de documentos, no dentro del proyecto de JARVIS."""
    nombre = os.path.basename((nombre or "").strip().strip('"')) or "documento.txt"
    return ruta_libre(carpeta_base(), nombre)


# ══════════════════════════════════════════════════════════════════════════
# ABRIR, PEGAR Y LA VENTANA CON EL FOCO
# ══════════════════════════════════════════════════════════════════════════
_TEXTO_PLANO = {"txt", "md", "csv", "json", "xml", "yml", "log", "ini", "cfg", "rtf", "html",
                "css", "sql", "tex"}


def abrir(ruta: str, app: str = "", log=print) -> tuple:
    """Abre el documento. Devuelve (abierto, con_que). Nunca ejecuta código."""
    if (os.getenv("JARVIS_REDACTOR_ABRIR", "1") or "1").strip() == "0":
        return False, "desactivado"
    if not os.path.isfile(ruta):
        return False, "el archivo no existe"
    ext = os.path.splitext(ruta)[1].lower().lstrip(".")
    app = (app or "").lower()
    es_codigo = ext in _LENGUAJES and ext != "html"
    if es_codigo and app not in ("notepad", "vscode"):
        app = "vscode" if _ruta_vscode() else "notepad"
    if app == "vscode" and not _ruta_vscode():
        app = "notepad"
    if app == "notepad" and ext in ("docx", "xlsx", "pptx", "pdf"):
        app = ""     # el Bloc de notas no sabe abrir eso
    if not app:
        app = _APP_DE_FORMATO.get(ext, "")

    if os.name != "nt":
        visor = shutil.which("xdg-open") or shutil.which("open")
        if not visor or es_codigo:
            return False, "en este sistema no sé abrirlo"
        try:
            subprocess.Popen([visor, ruta], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            return True, app
        except Exception as e:
            return False, str(e)

    import ejecutor
    comando = None
    if app == "notepad":
        comando = ["notepad.exe", ruta]
    elif app == "vscode":
        comando = [_ruta_vscode(), ruta]
    elif app in _EXES:
        exe = _ruta_app(_EXES[app])
        if exe:
            comando = [exe, ruta]
    if comando:
        ok, error, _ = ejecutor.lanzar(comando, origen="redactor", orden=os.path.basename(ruta),
                                       verificar_ms=300, log=log)
        if ok:
            return True, app
        log(f"[REDACTOR] No pude abrirlo con {app}: {error}")
    if es_codigo:
        ok, error, _ = ejecutor.lanzar(["notepad.exe", ruta], origen="redactor",
                                       orden=os.path.basename(ruta), verificar_ms=300, log=log)
        return (True, "notepad") if ok else (False, error)
    try:
        os.startfile(ruta)          # el programa que el señor tenga para ese tipo
        # Sin decir «Word» si a lo mejor lo abrió otro (LibreOffice, WordPad…).
        return True, ("navegador" if ext in ("html", "htm") else "")
    except OSError as e:
        if ext in _TEXTO_PLANO:
            ok, error, _ = ejecutor.lanzar(["notepad.exe", ruta], origen="redactor",
                                           orden=os.path.basename(ruta), verificar_ms=300, log=log)
            if ok:
                return True, "notepad"
        return False, f"no hay ningún programa para abrir los .{ext} ({e})"


_RX_VENTANA_PROPIA = re.compile(r"\bj\.?\s?a\.?\s?r\.?\s?v\.?\s?i\.?\s?s\b|\bultron\b", re.I)


def ventana_activa() -> str:
    """Título de la ventana con el foco ('' fuera de Windows o si no hay)."""
    if os.name != "nt":
        return ""
    try:
        import ctypes
        user32 = ctypes.windll.user32
        hwnd = user32.GetForegroundWindow()
        if not hwnd:
            return ""
        buf = ctypes.create_unicode_buffer(512)
        user32.GetWindowTextW(hwnd, buf, 512)
        return buf.value or ""
    except Exception:
        return ""


def es_ventana_de_jarvis(titulo: str) -> bool:
    """¿Es la propia interfaz de JARVIS/ULTRON (o su consola)?"""
    return bool(titulo) and bool(_RX_VENTANA_PROPIA.search(titulo))


def _al_portapapeles(texto: str) -> bool:
    try:
        import pyperclip
        pyperclip.copy(texto)
        return True
    except Exception:
        pass
    if os.name != "nt":
        return False
    try:
        tmp = os.path.join(tempfile.gettempdir(), "jarvis_pegar.txt")
        with open(tmp, "w", encoding="utf-8") as f:
            f.write(texto)
        r = subprocess.run(["powershell", "-NoProfile", "-Command",
                            f"Get-Content -LiteralPath '{tmp}' -Raw -Encoding UTF8 | Set-Clipboard"],
                           capture_output=True, timeout=20, creationflags=_SIN_VENTANA)
        return r.returncode == 0
    except Exception:
        return False


def pegar_en_ventana(texto: str, log=print) -> tuple:
    """Escribe el texto en la ventana con el foco. (ok, detalle).

    Por el portapapeles y Ctrl+V, no tecla a tecla: `pyautogui.typewrite` se
    salta las tildes, las eñes y los signos de apertura, que es como «Mañana»
    acababa siendo «Maana». Y nunca en la ventana del propio JARVIS.
    """
    titulo = ventana_activa()
    if not titulo:
        return False, "no hay ninguna ventana con el foco"
    if es_ventana_de_jarvis(titulo):
        return False, "la ventana con el foco es la mía"
    previo = None
    try:
        import pyperclip
        previo = pyperclip.paste()
    except Exception:
        pass
    if not _al_portapapeles(texto):
        return False, "no pude usar el portapapeles"
    try:
        import pyautogui
        time.sleep(0.25)
        pyautogui.hotkey("ctrl", "v")
    except Exception as e:
        return False, f"no pude pegar ({e})"
    if previo is not None:
        def _devolver():
            time.sleep(1.2)
            _al_portapapeles(previo)
        threading.Thread(target=_devolver, daemon=True).start()
    log(f"[REDACTOR] Texto pegado en «{titulo[:60]}» ({len(texto)} caracteres)")
    return True, titulo


# ══════════════════════════════════════════════════════════════════════════
# EL ENCARGO ENTERO
# ══════════════════════════════════════════════════════════════════════════
def _trato(core) -> str:
    return ", señor" if getattr(core, "nombre_agente", "JARVIS") == "JARVIS" else ""


_FEMENINOS = {"carta", "lista", "cancion", "presentacion", "tabla", "hoja", "pagina",
              "propuesta", "guia", "monografia", "tesis", "tesina", "nota", "receta",
              "solicitud", "cronica", "resena", "critica", "biografia", "invitacion",
              "explicacion", "historia", "factura", "cotizacion", "nomina", "acta", "frase",
              "adivinanza", "poesia", "novela", "leyenda", "fabula", "obra", "dedicatoria",
              "felicitacion", "publicacion", "descripcion", "definicion", "redaccion",
              "composicion", "sintesis", "planificacion", "encuesta", "ficha", "rutina",
              "dieta", "letra", "idea", "razon", "pregunta", "instruccion", "diapositiva"}
_SIN_PLURAL = {"tesis", "sintesis", "analisis", "brindis", "trabalenguas", "crisis"}


def _articulo(palabra: str) -> str:
    """«el», «la», «los» o «las» para la primera palabra de lo que se pide."""
    p = _plano((palabra or "").split()[0] if (palabra or "").split() else "")
    plural = p.endswith("s") and p not in _SIN_PLURAL
    singular = p[:-2] if p.endswith("es") and p[:-2] in _FEMENINOS else (p[:-1] if plural else p)
    femenino = singular in _FEMENINOS or singular.endswith(("cion", "sion", "dad", "tud"))
    if plural:
        return "las" if femenino else "los"
    return "la" if femenino else "el"


def _vista_previa(md: str, fam: str, datos=None) -> str:
    if fam == "texto":
        texto = re.sub(r"\s+", " ", a_texto_plano(md).replace("=", " ").replace("-", " ")).strip()
        titulo = titulo_de(md)
        if titulo and texto.startswith(titulo):
            texto = texto[len(titulo):].strip()
        return (texto[:220].rsplit(" ", 1)[0] + "…") if len(texto) > 220 else texto
    return ""


def _contar_palabras(md: str) -> int:
    return len(re.findall(r"\w+", _sin_marcas(md or "")))


def _registrar(core, ruta: str, enc: dict, log):
    try:
        import deshacer
        deshacer.anotar("archivo", f"creé {os.path.basename(ruta)}",
                        {"ruta": ruta, "respaldo": None, "existia": False},
                        agente=getattr(core, "nombre_agente", "JARVIS"), log=log)
    except Exception:
        pass
    try:
        guardar = getattr(core, "save_media_history", None)
        if callable(guardar):
            guardar("documento", enc.get("encargo", "")[:300], ruta)
    except Exception:
        pass


def _aviso_previo(enc: dict, formato: str, app: str, tr: str) -> str:
    """Lo que se dice ANTES de ponerse a escribir: redactar un ensayo con el
    cerebro de casa lleva su rato y el silencio parece un cuelgue."""
    tipo = enc.get("tipo") or "documento"
    descripcion = enc.get("descripcion") or ""
    palabras = descripcion.split()
    primera = _plano(palabras[0]) if palabras else ""
    if primera == "algo":
        que, articulo = descripcion, "el"
    elif primera in _CANTIDADES or primera.isdigit():
        # «10 chistes», «unas ideas»: sin artículo; el pronombre, del sustantivo.
        que = descripcion
        articulo = _articulo(palabras[1]) if len(palabras) > 1 else "el"
        if primera not in ("un", "una") and articulo in ("el", "la"):
            articulo = "los" if articulo == "el" else "las"
    elif descripcion:
        articulo = _articulo(descripcion)
        que = f"{articulo} {descripcion}"
    else:
        articulo = _articulo(tipo)
        que = f"{articulo} {tipo}" + (f" sobre {enc['tema']}" if enc.get("tema") else "")
    fam = familia(formato)
    verbo = "preparo" if fam in ("tabla", "diapos") else "escribo"
    destino = _NOMBRE_APP.get(app) or {"docx": "Word", "xlsx": "Excel", "pptx": "PowerPoint",
                                       "txt": "el Bloc de notas"}.get(formato, "")
    if enc.get("pegar"):
        return f"Enseguida{tr}. Le {verbo} {que} y lo escribo donde tiene el cursor."
    pronombre = {"el": "lo", "la": "la", "los": "los", "las": "las"}.get(articulo, "lo")
    cola = f" y se {pronombre} abro en {destino}." if destino and fam != "codigo" else "."
    return f"Enseguida{tr}. Le {verbo} {que}{cola}"


def crear(core, enc: dict, log=print, avisar=None, abrir_al_final: bool = True,
          carpeta_destino: str = "") -> dict:
    """Hace el encargo entero. Devuelve {ok, mensaje, voz, ruta, ...}."""
    tr = _trato(core)
    formato = (enc.get("formato") or "txt").lower().lstrip(".")
    fam = familia(formato)
    app = enc.get("app") or ""
    tipo = enc.get("tipo") or "documento"

    # 1. El contenido
    if enc.get("modo") == "literal":
        contenido = enc.get("contenido") or ""
        if fam == "tabla":
            filas = [re.split(r"\s*[;,\t]\s*", l) for l in contenido.splitlines() if l.strip()]
            contenido = {"titulo": enc.get("tema") or "Tabla",
                         "hojas": [{"columnas": filas[0] if filas else ["Texto"],
                                    "filas": filas[1:]}]}
    elif enc.get("modo") == "anterior":
        previa = enc.get("contenido") or _ultima_respuesta(core)
        if not previa:
            return {"ok": False, "mensaje": f"No tengo nada reciente que guardar{tr}. "
                                            "Dígame qué quiere que le escriba.",
                    "voz": ""}
        if fam == "texto":
            contenido = previa
        else:
            # A Excel, PowerPoint o código no se pega la respuesta tal cual:
            # se le pide al cerebro que la convierta.
            if avisar:
                avisar(f"Enseguida{tr}. Lo paso a {_NOMBRE_APP.get(app) or formato.upper()}.")
            contenido = _redactar(core, dict(enc, encargo=f"Convierte este texto en "
                                             f"{tipo}: {enc['encargo']}"), log=log,
                                  texto_base=previa)
            if contenido is None:
                return {"ok": False, "voz": "",
                        "mensaje": f"No he podido convertirlo{tr}: {ultimo_fallo() or 'el cerebro no contestó'}."}
    else:
        if avisar:
            avisar(_aviso_previo(enc, formato, app, tr))
        base = _ultima_respuesta(core) if enc.get("usar_anterior") else ""
        contenido = _redactar(core, enc, log=log, texto_base=base)
        if contenido is None:
            motivo = ultimo_fallo() or "el cerebro no devolvió nada aprovechable"
            return {"ok": False, "voz": f"No he podido escribirlo{tr}.",
                    "mensaje": (f"No he podido escribirlo{tr}: {motivo}. No le dejo un "
                                "documento con plantillas vacías; en cuanto vuelva el cerebro, "
                                "pídamelo otra vez.")}

    # 2. «Escríbelo aquí»: en la ventana con el foco, si no es la mía.
    if enc.get("pegar"):
        if isinstance(contenido, dict) and fam == "tabla":
            # Separado por tabuladores: pegado en Excel o en Sheets cae en celdas.
            texto = "\n".join("\t".join("" if v is None else str(v) for v in fila)
                              for h in _hojas(contenido) for fila in [h["columnas"]] + h["filas"])
        elif isinstance(contenido, dict):
            texto = a_texto_plano(_diapos_a_md(contenido)).strip()
        elif fam == "texto" and enc.get("modo") != "literal":
            texto = a_texto_plano(contenido).strip()
        else:
            texto = str(contenido)
        ok, detalle = pegar_en_ventana(texto, log=log)
        if ok:
            msg = f"Hecho{tr}. Lo he escrito en «{detalle[:60]}»."
            _recordar({"ruta": "", "formato": formato, "app": "", "tipo": tipo, "ts": time.time(),
                       "titulo": "", "contenido": texto, "encargo": enc.get("encargo", ""),
                       "respuesta": msg})
            return {"ok": True, "mensaje": msg, "voz": msg, "ruta": ""}
        log(f"[REDACTOR] No lo pego ({detalle}); va al Bloc de notas")
        if isinstance(contenido, dict):
            contenido = _tabla_a_md(contenido) if fam == "tabla" else _diapos_a_md(contenido)
        app, formato, fam = "notepad", "txt", "texto"

    # 3. Nombre y carpeta
    if isinstance(contenido, dict):
        titulo = str(contenido.get("titulo") or "").strip()
    elif fam == "texto":
        titulo = titulo_de(contenido)
    else:
        titulo = ""
    nombre = enc.get("nombre") or ""
    if nombre and not os.path.splitext(nombre)[1]:
        nombre = f"{nombre}.{formato}"
    if nombre:
        nombre = nombre_archivo(os.path.splitext(nombre)[0], formato)
    elif enc.get("modo") in ("literal", "anterior") and isinstance(contenido, str) and not titulo:
        # Lo dictado o lo de antes no trae título: las primeras palabras sirven
        # más que «Documento (7).txt».
        primeras = " ".join(_sin_marcas(contenido).split()[:7])
        nombre = nombre_archivo(re.split(r"[.!?;:\n]", primeras)[0] or primeras, formato)
    elif fam == "codigo":
        nombre = nombre_codigo(enc.get("descripcion") or enc.get("tema") or "", formato)
    else:
        nombre = nombre_archivo(titulo or _mayuscula(enc.get("tema") or "") or _mayuscula(tipo),
                                formato)
    destino = carpeta_destino or enc.get("ubicacion") or carpeta_base()
    ext_destino = os.path.splitext(destino)[1].lower().lstrip(".")
    if ext_destino and (ext_destino in _FAMILIA or ext_destino in _LENGUAJES):
        # «guárdalo en C:\\Users\\X\\Desktop\\carta.docx»: eso es el archivo, no la carpeta.
        destino, nombre = os.path.dirname(destino) or carpeta_base(), os.path.basename(destino)
        if ext_destino != formato:
            nombre = os.path.splitext(nombre)[0] + "." + formato
    try:
        os.makedirs(destino, exist_ok=True)
    except OSError as e:
        return {"ok": False, "voz": "", "mensaje": f"No puedo usar la carpeta {destino}{tr}: {e}"}
    ruta = ruta_libre(destino, nombre)

    # 4. Escribir y comprobar
    try:
        ruta, aviso = escribir(ruta, formato, contenido, titulo=titulo,
                               literal=enc.get("modo") == "literal", log=log)
    except PermissionError:
        return {"ok": False, "voz": "", "mensaje": f"No tengo permiso para escribir en {destino}{tr}."}
    except Exception as e:
        log(f"[REDACTOR] No pude escribir {ruta}: {type(e).__name__}: {e}")
        return {"ok": False, "voz": "",
                "mensaje": f"No he podido escribir el archivo{tr}: {str(e)[:120]}"}
    if not os.path.isfile(ruta) or os.path.getsize(ruta) == 0:
        return {"ok": False, "voz": "", "mensaje": f"Escribí el archivo pero no lo encuentro en disco{tr}."}
    formato = os.path.splitext(ruta)[1].lower().lstrip(".")
    fam = familia(formato)
    _registrar(core, ruta, enc, log)
    log(f"[REDACTOR] Documento listo: {ruta} ({os.path.getsize(ruta)} bytes)")

    # 5. Abrirlo con el texto dentro
    abierto, con = (False, "")
    if abrir_al_final:
        if app and _FORMATO_DE_APP.get(app) and familia(_FORMATO_DE_APP[app]) != fam and \
                not (app == "notepad" and fam in ("texto", "codigo")):
            app = ""
        abierto, con = abrir(ruta, app, log=log)

    # 6. Lo que se le dice al señor
    md_texto = contenido if isinstance(contenido, str) else ""
    if fam == "tabla" and isinstance(contenido, dict):
        hojas = _hojas(contenido)
        n_filas = sum(len(h["filas"]) for h in hojas)
        que = (f"la hoja «{titulo or os.path.splitext(os.path.basename(ruta))[0]}», con {n_filas} "
               f"filas" + (f" en {len(hojas)} hojas" if len(hojas) > 1 else ""))
    elif fam == "diapos" and isinstance(contenido, dict):
        n = len(contenido.get("diapositivas") or [])
        que = f"la presentación «{titulo or 'sin título'}», de {n + 1} diapositivas"
    elif fam == "codigo":
        que = f"«{os.path.basename(ruta)}» ({len(str(contenido).splitlines())} líneas)"
    elif enc.get("modo") == "literal":
        que = "su texto"
    elif enc.get("modo") == "anterior":
        que = "lo que le acababa de decir"
    else:
        n = _contar_palabras(md_texto)
        que = f"«{titulo or nombre.rsplit('.', 1)[0]}»" + (f" ({n} palabras)" if n > 40 else "")
    if enc.get("modo") == "literal":
        verbo = "He escrito"
    elif enc.get("modo") == "anterior":
        verbo = "He guardado"
    else:
        verbo = "Le he escrito" if fam in ("texto", "codigo") else "Le he preparado"
    nombre_app = _NOMBRE_APP.get(con) or ("su programa de siempre" if abierto else "")
    if abierto:
        pronombre = "la" if fam in ("tabla", "diapos") else "lo"
        cabeza = f"Listo{tr}. {verbo} {que} y se {pronombre} he abierto en {nombre_app}."
    elif not abrir_al_final or con == "desactivado":
        cabeza = f"Listo{tr}. {verbo} {que}."
    else:
        cabeza = f"Listo{tr}. {verbo} {que}, pero no he podido abrirlo: {con}."
    lineas = [cabeza, f"Está en {ruta}"]
    if fam == "codigo":
        error = _revisar_codigo(formato, str(contenido))
        if error:
            lineas.append(f"Ojo: no compila del todo ({error}). Revíselo antes de usarlo.")
        else:
            lineas.append("No lo he ejecutado: eso lo decide usted.")
    if aviso:
        lineas.append(aviso)
    vista = _vista_previa(md_texto, fam) if enc.get("modo") == "generar" else ""
    if vista:
        lineas.append(f"Empieza así: «{vista}»")
    mensaje = "\n\n".join(lineas)

    _recordar({"ruta": ruta, "formato": formato, "app": con if abierto else app, "tipo": tipo,
               "titulo": titulo, "encargo": enc.get("encargo", ""), "ts": time.time(),
               "contenido": json.dumps(contenido, ensure_ascii=False)
               if isinstance(contenido, dict) else str(contenido),
               "respuesta": mensaje})
    return {"ok": True, "mensaje": mensaje, "voz": cabeza, "ruta": ruta, "formato": formato,
            "titulo": titulo, "abierto": abierto, "app": con or app, "contenido": contenido}


# ══════════════════════════════════════════════════════════════════════════
# SEGUIMIENTOS DEL ÚLTIMO DOCUMENTO
# ══════════════════════════════════════════════════════════════════════════
def _contenido_ultimo(ult: dict):
    crudo = ult.get("contenido") or ""
    fam = familia(ult.get("formato", ""))
    if fam in ("tabla", "diapos"):
        try:
            return json.loads(crudo) if isinstance(crudo, str) else crudo
        except Exception:
            return None
    return crudo or leer_contenido(ult["ruta"], ult.get("formato", ""))


def enviar_al_movil(ruta: str, log=print) -> str:
    """Copia el documento a /envios del servidor web y devuelve el enlace."""
    from urllib.parse import quote
    base = os.path.join(os.path.dirname(os.path.abspath(__file__)), "web_interface", "envios")
    os.makedirs(base, exist_ok=True)
    nombre = os.path.basename(ruta)
    destino = os.path.join(base, nombre)
    if os.path.exists(destino):
        nombre = time.strftime("%H%M%S_") + nombre
        destino = os.path.join(base, nombre)
    shutil.copy2(ruta, destino)
    import jarvis_config
    url = f"http://{jarvis_config.LOCAL_IP}:{jarvis_config.PORT}/envios/{quote(nombre)}"
    try:
        import urllib.request
        req = urllib.request.Request(jarvis_config.url_flask("/notify"),
                                     data=json.dumps({"text": f"Documento listo: {url}"}).encode("utf-8"),
                                     headers={"Content-Type": "application/json"})
        urllib.request.urlopen(req, timeout=1.5)
    except Exception:
        pass
    return url


def _seguimiento(core, enc: dict, log=print, avisar=None) -> dict:
    ult = enc["ultimo"]
    tr = _trato(core)
    ruta = ult["ruta"]
    accion = enc.get("accion")
    nombre = os.path.basename(ruta)

    if accion == "donde":
        msg = f"«{nombre}» está en {os.path.dirname(ruta)}{tr}."
        return {"ok": True, "mensaje": msg, "voz": f"Está en la carpeta {os.path.basename(os.path.dirname(ruta))}{tr}."}
    if accion == "abrir":
        # «Ábrelo con Word»: manda el programa que diga ahora.
        pedido, _f = _detectar_destino(_plano(enc.get("frase", "")))
        ok, con = abrir(ruta, pedido or ult.get("app", ""), log=log)
        msg = (f"Abierto{tr}: «{nombre}» en {_NOMBRE_APP.get(con) or 'su programa de siempre'}."
               if ok else f"No he podido abrir «{nombre}»{tr}: {con}.")
        return {"ok": ok, "mensaje": msg, "voz": msg}
    if accion == "leer":
        texto = a_texto_plano(_contenido_ultimo(ult) or "") if familia(ult.get("formato")) == "texto" else ""
        if not texto.strip():
            return {"ok": False, "mensaje": f"Ese documento no lo puedo leer en voz alta{tr}.", "voz": ""}
        return {"ok": True, "mensaje": texto.strip(), "voz": texto.strip()[:4000]}
    if accion == "enviar":
        try:
            url = enviar_al_movil(ruta, log=log)
        except Exception as e:
            return {"ok": False, "voz": "", "mensaje": f"No he podido mandarlo al móvil{tr}: {e}"}
        msg = f"Listo{tr}: «{nombre}» está en su teléfono.\n\n{url}"
        return {"ok": True, "mensaje": msg, "voz": f"Listo{tr}: se lo he mandado al teléfono."}

    contenido = _contenido_ultimo(ult)
    if contenido in (None, ""):
        return {"ok": False, "voz": "",
                "mensaje": f"No puedo leer «{nombre}» para trabajar sobre él{tr}."}

    if accion == "convertir":
        nuevo = enc["formato"]
        fam_vieja, fam_nueva = familia(ult.get("formato")), familia(nuevo)
        datos = contenido
        if fam_vieja != fam_nueva:
            texto_base = contenido if isinstance(contenido, str) else (
                _tabla_a_md(contenido) if fam_vieja == "tabla" else _diapos_a_md(contenido))
            if fam_nueva == "texto":
                datos = texto_base
            else:
                if avisar:
                    avisar(f"Enseguida{tr}. Lo paso a {nuevo.upper()}.")
                datos = _redactar(core, {"formato": nuevo, "encargo": f"Convierte este documento: "
                                         f"{ult.get('encargo', '')}", "tipo": ult.get("tipo", "")},
                                  log=log, texto_base=texto_base)
                if datos is None:
                    return {"ok": False, "voz": "",
                            "mensaje": f"No he podido convertirlo{tr}: {ultimo_fallo()}."}
        destino = ruta_libre(os.path.dirname(ruta), os.path.splitext(nombre)[0] + "." + nuevo)
        try:
            destino, aviso = escribir(destino, nuevo, datos, titulo=ult.get("titulo", ""), log=log)
        except Exception as e:
            return {"ok": False, "voz": "", "mensaje": f"No he podido convertirlo{tr}: {e}"}
        _registrar(core, destino, {"encargo": ult.get("encargo", "")}, log)
        ok, con = abrir(destino, "", log=log)
        cabeza = (f"Listo{tr}. Se lo he pasado a {nuevo.upper()}"
                  + (f" y lo tiene abierto en {_NOMBRE_APP.get(con) or 'su programa de siempre'}." if ok else "."))
        msg = cabeza + f"\n\nEstá en {destino}" + (f"\n\n{aviso}" if aviso else "")
        _recordar(dict(ult, ruta=destino, formato=os.path.splitext(destino)[1].lstrip("."),
                       app=con, ts=time.time(), respuesta=msg,
                       contenido=json.dumps(datos, ensure_ascii=False) if isinstance(datos, dict)
                       else str(datos)))
        return {"ok": True, "mensaje": msg, "voz": cabeza, "ruta": destino}

    # modificar
    fam = familia(ult.get("formato"))
    if avisar:
        avisar(f"Enseguida{tr}. Me pongo con el cambio.")
    base_txt = json.dumps(contenido, ensure_ascii=False, indent=1) if isinstance(contenido, dict) \
        else str(contenido)
    usuario = (f"Documento actual ({ult.get('tipo') or 'documento'}, formato "
               f"{ult.get('formato')}):\n<<<\n{base_txt[:30000]}\n>>>\n\n"
               f"Cambio que pide el señor: «{enc['frase']}»")
    tope = max(2000, min(int(_contar_palabras(base_txt) * 2.6) + 1500, 16000))
    crudo = preguntar(core, _SISTEMA_CAMBIO, usuario, tope=tope, temperatura=0.4, log=log)
    if fam in ("tabla", "diapos"):
        nuevo = _rescatar_json(crudo)
        if not isinstance(nuevo, dict):
            nuevo = None
    elif fam == "codigo":
        nuevo = _limpiar_codigo(crudo) if crudo.strip() else None
    else:
        nuevo = _quitar_preambulo(crudo) or None
    if not nuevo:
        return {"ok": False, "voz": f"No he podido hacer el cambio{tr}.",
                "mensaje": f"No he podido hacer el cambio{tr}: {ultimo_fallo() or 'el cerebro no contestó'}."}
    # Respaldo para «deshaz eso» y, si Word lo tiene bloqueado, versión nueva.
    respaldo = None
    try:
        import herramientas_fs
        respaldo = herramientas_fs._respaldar(ruta, log=log)
    except Exception:
        pass
    destino = ruta
    try:
        destino, aviso = escribir(ruta, ult.get("formato"), nuevo, titulo=ult.get("titulo", ""), log=log)
    except PermissionError:
        destino = ruta_libre(os.path.dirname(ruta), nombre)
        destino, aviso = escribir(destino, ult.get("formato"), nuevo, titulo=ult.get("titulo", ""), log=log)
        aviso = ("El original lo tiene abierto otro programa, así que lo he guardado como "
                 "versión nueva. " + aviso).strip()
    except Exception as e:
        return {"ok": False, "voz": "", "mensaje": f"No he podido guardar el cambio{tr}: {e}"}
    try:
        import deshacer
        deshacer.anotar("archivo", f"cambié {os.path.basename(destino)}",
                        {"ruta": destino, "respaldo": respaldo if destino == ruta else None,
                         "existia": destino == ruta},
                        agente=getattr(core, "nombre_agente", "JARVIS"), log=log)
    except Exception:
        pass
    ok, con = abrir(destino, ult.get("app", ""), log=log)
    titulo = (nuevo.get("titulo") if isinstance(nuevo, dict) else titulo_de(nuevo)) or ult.get("titulo", "")
    cabeza = f"Hecho{tr}. He cambiado «{titulo or nombre}»" + (
        f" y lo tiene abierto en {_NOMBRE_APP.get(con) or 'su programa de siempre'}." if ok else ".")
    msg = cabeza + f"\n\nEstá en {destino}" + (f"\n\n{aviso}" if aviso else "")
    _recordar(dict(ult, ruta=destino, ts=time.time(), titulo=titulo, respuesta=msg,
                   contenido=json.dumps(nuevo, ensure_ascii=False) if isinstance(nuevo, dict)
                   else str(nuevo)))
    return {"ok": True, "mensaje": msg, "voz": cabeza, "ruta": destino}


# ══════════════════════════════════════════════════════════════════════════
# PUERTAS DE ENTRADA
# ══════════════════════════════════════════════════════════════════════════
def atender_detalle(core, texto: str, log=print, avisar=None):
    """Para el núcleo: {mensaje, voz} si la frase era un encargo de escribir,
    o None si no lo era (y entonces sigue el camino de siempre)."""
    enc = entender(texto, core=core)
    if not enc:
        return None
    log(f"[REDACTOR] Encargo: {enc.get('modo')} {enc.get('accion', '')} "
        f"{enc.get('tipo', '')} -> {enc.get('formato', '')} {enc.get('app', '')}".strip())
    if enc["modo"] == "seguimiento":
        r = _seguimiento(core, enc, log=log, avisar=avisar)
    else:
        r = crear(core, enc, log=log, avisar=avisar)
    return {"mensaje": r["mensaje"], "voz": r.get("voz", "")}


def atender(core, texto: str, log=print, avisar=None):
    """Como atender_detalle, pero solo el texto de la respuesta (o None)."""
    r = atender_detalle(core, texto, log=log, avisar=avisar)
    return r["mensaje"] if r else None


def crear_documento(core, encargo: str = "", formato: str = "", nombre: str = "",
                    ubicacion: str = "", abrir_con: str = "", contenido: str = "",
                    log=print) -> str:
    """La herramienta del bucle del modelo («crear_documento»).

    El modelo ya ha decidido que hay que escribir algo: aquí no hace falta
    adivinar la intención, solo completar lo que no diga con lo que se deduzca
    del propio encargo.
    """
    encargo = (encargo or "").strip()
    pista = entender(encargo) if encargo else None
    if pista and pista.get("modo") == "seguimiento":
        pista = None
    formato = (formato or "").lower().strip().lstrip(".")
    formato = {"word": "docx", "excel": "xlsx", "powerpoint": "pptx", "texto": "txt",
               "bloc de notas": "txt", "markdown": "md", "python": "py",
               "javascript": "js"}.get(formato, formato)
    app = (abrir_con or "").lower().strip()
    app = {"bloc de notas": "notepad", "notas": "notepad", "microsoft word": "word",
           "power point": "powerpoint", "visual studio code": "vscode", "code": "vscode",
           "chrome": "navegador", "edge": "navegador", "navegador web": "navegador"}.get(app, app)
    if app and app not in _NOMBRE_APP:
        app = ""
    if not formato:
        formato = (pista or {}).get("formato") or (_FORMATO_DE_APP.get(app) if app else "") or \
            os.path.splitext(nombre or "")[1].lstrip(".").lower() or "txt"
    destino = ""
    if ubicacion:
        destino = carpeta(ubicacion) or (os.path.abspath(os.path.expanduser(ubicacion))
                                         if os.path.isabs(os.path.expanduser(ubicacion)) else
                                         os.path.join(carpeta_base(), ubicacion))
    enc = {
        "modo": "literal" if contenido else "generar",
        "frase": encargo, "encargo": encargo or (nombre or "documento"),
        "contenido": contenido or "",
        "tipo": (pista or {}).get("tipo") or ("documento" if familia(formato) == "texto"
                                              else familia(formato)),
        "formato": formato, "app": app or (pista or {}).get("app", ""),
        "ubicacion": destino or (pista or {}).get("ubicacion", ""),
        "nombre": nombre or (pista or {}).get("nombre", ""),
        "tema": (pista or {}).get("tema", ""),
        "palabras": (pista or {}).get("palabras") or _palabras(_plano(encargo), "", familia(formato)),
        "pegar": False,
    }
    if contenido and familia(formato) in ("tabla", "diapos"):
        datos = _rescatar_json(contenido)
        if isinstance(datos, dict):
            enc["contenido"] = ""
            r = _escribir_directo(core, enc, datos, log)
            return r["mensaje"]
    if not contenido and not encargo:
        return "Falta el encargo: diga qué hay que escribir."
    return crear(core, enc, log=log)["mensaje"]


def _escribir_directo(core, enc, datos, log):
    """Una tabla o una presentación que el modelo ya trae hecha en JSON."""
    titulo = str(datos.get("titulo") or "").strip()
    destino = enc.get("ubicacion") or carpeta_base()
    os.makedirs(destino, exist_ok=True)
    ruta = ruta_libre(destino, nombre_archivo(enc.get("nombre") or titulo or enc.get("tipo"),
                                              enc["formato"]))
    ruta, aviso = escribir(ruta, enc["formato"], datos, titulo=titulo, log=log)
    _registrar(core, ruta, enc, log)
    ok, con = abrir(ruta, enc.get("app", ""), log=log)
    msg = (f"Listo{_trato(core)}. Está en {ruta}" + (f" y abierto en {_NOMBRE_APP.get(con) or 'su programa de siempre'}." if ok else ".")
           + (f" {aviso}" if aviso else ""))
    _recordar({"ruta": ruta, "formato": enc["formato"], "app": con, "tipo": enc.get("tipo", ""),
               "titulo": titulo, "encargo": enc.get("encargo", ""), "ts": time.time(),
               "contenido": json.dumps(datos, ensure_ascii=False), "respuesta": msg})
    return {"ok": True, "mensaje": msg, "ruta": ruta}


def generar_archivo(core, encargo: str, formato: str, carpeta_destino: str = "",
                    abrir_al_final: bool = False, log=print) -> dict:
    """Para jarvis_generator y la web: el archivo sin abrir nada."""
    pista = entender(encargo) or {}
    if pista.get("modo") == "seguimiento":
        pista = {}
    enc = {"modo": "generar", "frase": encargo, "encargo": encargo, "contenido": "",
           "tipo": pista.get("tipo") or ("documento" if familia(formato) == "texto"
                                         else familia(formato)),
           "formato": formato, "app": "", "ubicacion": "", "nombre": "",
           "tema": pista.get("tema", ""), "pegar": False,
           "palabras": pista.get("palabras") or _palabras(_plano(encargo), pista.get("tipo", ""),
                                                          familia(formato))}
    return crear(core, enc, log=log, abrir_al_final=abrir_al_final,
                 carpeta_destino=carpeta_destino)


def pide_escribir(texto: str) -> bool:
    """¿La frase pide escribir o poner algo? Más laxo que `entender`: es la red
    del chat, para cuando el encargo llegó hasta el modelo."""
    p = _plano(texto or "")
    return bool(_RX_V_ESCRIBIR.search(p) or _RX_V_PONER.search(p)
                or (_RX_V_CREAR.search(p) and _detectar_genero(p))) \
        and not _RX_MENSAJERIA.search(p)


def app_de_escritura(nombre: str) -> str:
    """'notepad', 'word' o 'wordpad' si es un programa donde se escribe; si no, ''."""
    p = _plano(nombre or "").strip()
    if p in ("notepad.exe", "notepad", "notas"):
        return "notepad"
    if p in ("winword", "winword.exe"):
        return "word"
    for rx, app in _PISTAS_APP:
        if rx.search(p) and app in ("notepad", "word", "wordpad"):
            return app
    return ""


def guardar_respuesta(core, texto: str, app_nombre: str, log=print) -> str:
    """El modelo contestó con [OPEN:bloc de notas] y el texto en el chat: se
    guarda ese texto y el programa se abre CON él. Devuelve la nota para el
    señor, o '' si no había texto que guardar (y entonces se abre sin más)."""
    app = app_de_escritura(app_nombre) or "notepad"
    limpio = re.sub(r"\[OPEN:[^\]]+\]", "", texto or "").strip()
    limpio = _quitar_preambulo(limpio)
    # Lo que queda tras quitar la cortesía del principio: «Abriendo el Bloc de
    # notas, señor.», «Claro, señor. Aquí tiene su lista:»…
    cuerpo = limpio
    for _ in range(4):
        nuevo = re.sub(r"^\W*(?:(?:¡?claro|por supuesto|desde luego|enseguida|con (?:mucho )?gusto"
                       r"|perfecto|listo|muy bien|de acuerdo|abriendo|abro|le abro|ya tiene)\b"
                       r"[^.!:\n]*[.!:]|aqu[ií] (?:tiene|tienes|est[aá]|va|le dejo)\b[^:\n]*:)\s*",
                       "", cuerpo, flags=re.I).strip()
        if nuevo == cuerpo:
            break
        cuerpo = nuevo
    if len(cuerpo) < 40:
        abrir_vacio = getattr(core, "_open_app", None)
        if callable(abrir_vacio):
            abrir_vacio(app_nombre)
        return ""
    formato = _FORMATO_DE_APP.get(app, "txt")
    enc = {"modo": "literal", "frase": texto, "encargo": cuerpo[:200], "contenido": cuerpo,
           "tipo": "texto", "formato": formato, "app": app, "ubicacion": "", "nombre": "",
           "tema": "", "descripcion": "", "palabras": 0, "pegar": False}
    r = crear(core, enc, log=log)
    if not r.get("ok"):
        return ""
    if r.get("abierto"):
        return (f"Se lo he dejado escrito en {_NOMBRE_APP.get(r.get('app') or app, 'su programa')}: "
                f"{r['ruta']}")
    return f"Se lo he dejado escrito en {r['ruta']}"


def teclear(core, texto: str, log=print) -> str:
    """«teclea hola», «escribe buenos días» desde pc_control.

    Con tildes y eñes (portapapeles, no tecla a tecla) y NUNCA en la ventana
    del propio JARVIS: si el foco está ahí, el texto va a un Bloc de notas.
    """
    texto = (texto or "").strip()
    if not texto:
        return f"¿Qué quiere que escriba{_trato(core)}?"
    ok, detalle = pegar_en_ventana(texto, log=log)
    if ok:
        return f"Escrito{_trato(core)}: «{texto[:80]}» en «{detalle[:50]}»."
    log(f"[REDACTOR] No tecleo ({detalle}); lo dejo en el Bloc de notas")
    enc = {"modo": "literal", "frase": texto, "encargo": texto, "contenido": texto,
           "tipo": "texto", "formato": "txt", "app": "notepad", "ubicacion": "",
           "nombre": "", "tema": " ".join(texto.split()[:6]), "descripcion": "",
           "palabras": 0, "pegar": False}
    r = crear(core, enc, log=log)
    if not r.get("ok"):
        return r["mensaje"]
    porque = ("Para no escribir en mi propia ventana" if "mía" in detalle
              else "Como no había ninguna ventana donde escribir")
    donde = "se lo he dejado en el Bloc de notas" if r.get("abierto") else "se lo he guardado"
    return f"{porque}, {donde}{_trato(core)}.\n\nEstá en {r['ruta']}"


def texto_a_teclear(frase: str) -> str:
    """De «teclea Mañana voy», «escribe aquí: Hola», lo que hay que escribir,
    con sus mayúsculas y tildes (pc_control solo ve la frase normalizada)."""
    orig = (frase or "").strip()
    p = _plano(orig)
    m = re.search(r"\b(?:escribe|escriba|teclea|teclee)\b\s*(?:aqui|aca|en la ventana(?: activa)?"
                  r"|donde (?:esta|tengo) el cursor|literalmente|textualmente)?\s*[:,]?\s*", p)
    if not m:
        return ""
    # «escribe a Juan que llego tarde» es un mensaje: lo atiende mensajería.
    if _RX_MENSAJERIA.search(p) or re.match(r"(?:a|al|para)\s+\w+.*\b(?:que|diciendo)\b",
                                            p[m.end():]):
        return ""
    texto = orig[m.end():].strip()
    c = re.match(r"^[«\"“'](.*)[»\"”']$", texto, re.S)
    return (c.group(1) if c else texto).strip()


if __name__ == "__main__":
    try:
        import consola_utf8  # noqa: F401
    except Exception:
        pass
    for frase in sys.argv[1:] or ["abre el bloc de notas y escribe un poema sobre la luna"]:
        print(frase, "->", json.dumps(entender(frase), ensure_ascii=False, indent=1))
